from __future__ import annotations

import argparse
import base64
import csv
import concurrent.futures as futures
import getpass
import hashlib
import html
import http.client
import ipaddress
import io
import json
import os
import queue
import random
import re
import sqlite3
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable

from PIL import Image, ImageOps

from paddle_ocr import PaddleImage, PaddleOcrError, post_batch, verify_paddle_available


APP_VERSION = "0.7.0-paddle"
SQL_RULE_VERSION = "knime-db-rules-v1"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
MAX_PRODUCTS = 100
OCR_WORKERS = 3
OCR_MAX_WORKERS = 6
QWEN_EXTRACT_WORKERS = 3
QWEN_SEARCH_WORKERS = 2
OCR_MAX_ATTEMPTS = 5
QWEN_MAX_ATTEMPTS = 5
VLLM_NETWORK_FAILURE_THRESHOLD = 3
OCR_REQUEST_TIMEOUT = 30
OCR_FAST_NETWORK_FAILURE_SECONDS = 3.0
OCR_FAST_NETWORK_RETRY_DELAY = 2.0
OCR_RECOVERY_SUCCESS_COUNT = 3
OCR_IMAGE_MAX_LONG_EDGE = 1600
OCR_IMAGE_MAX_PIXELS = 2_000_000
OCR_IMAGE_JPEG_QUALITY = 90
OCR_RETRY_DELAYS = (5.0, 15.0, 30.0, 60.0)
PADDLE_OCR_PROMPT_VERSION = "paddle-ocr-payload-v1"
PADDLE_BATCH_SIZE = 8
PADDLE_DEFAULT_API_URL = "http://192.168.1.115:8870/v1/ocr"
PADDLE_DEFAULT_MODEL_VERSION = "PaddleOCR-VL-1.6"
PADDLE_RECOVERY_ATTEMPTS = 30
PADDLE_RECOVERY_DELAY_SECONDS = 10
QWEN_RETRY_DELAYS = (2.0, 5.0, 10.0, 20.0)
RETENTION_DAYS = 90
QWEN_BASE_URL = "https://dashscope.aliyuncs.com/compatible-mode/v1"
VLLM_OCR_PROMPT_VERSION = "qwen-vl-ocr-transcription-v2"
VLLM_OCR_SCHEMA = {
    "type": "object",
    "properties": {
        "blocks": {
            "type": "array",
            "minItems": 1,
            "items": {
                "type": "object",
                "properties": {
                    "order": {"type": "integer", "minimum": 1},
                    "text": {"type": "string", "minLength": 1},
                    "legibility": {"type": "string", "enum": ["clear", "unclear"]},
                },
                "required": ["order", "text", "legibility"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["blocks"],
    "additionalProperties": False,
}
VLLM_OCR_SYSTEM_PROMPT = """你是商品包装 OCR 转写员。请逐字转写图片中全部可见文字。

覆盖中文、英文、数字、标点、单位、配料、生产信息、地址、许可证、促销语和水印；
不得翻译、纠错、补全或推断。blocks 按阅读顺序编号，重复文字也必须保留。
难以辨认但仍可见的文字不能删除，legibility 标为 unclear。
只返回符合 JSON Schema 的对象。"""
PRODUCT_KEY_SEPARATOR = "\x1f"
ENRICHMENT_FIELDS = (
    "factory_name",
    "factory_address",
    "production_license",
    "origin",
)
ENRICHMENT_OUTPUT_NAMES = {
    "代工厂": "factory_name",
    "代工厂地址": "factory_address",
    "生产许可证": "production_license",
    "产地": "origin",
    "热门话题": "hot_topics",
}
ENRICHMENT_ATTRIBUTE_ALIASES = {
    "factory_name": (
        "代工厂",
        "生产企业",
        "生产商",
        "制造商",
        "生产厂家",
        "生产厂商",
        "企业名称",
        "厂名",
    ),
    "factory_address": (
        "代工厂地址",
        "生产地址",
        "厂址",
        "厂家地址",
        "企业地址",
    ),
    "production_license": (
        "生产许可证",
        "生产许可证编号",
        "食品生产许可证编号",
        "食品生产许可证",
    ),
    "origin": (
        "产地",
        "原产地",
        "原产国",
        "生产国",
        "生产地区",
        "进口国",
    ),
}


class DemoError(RuntimeError):
    def __init__(self, code: str, message: str, retryable: bool = False):
        super().__init__(message)
        self.code = code
        self.retryable = retryable
        self.attempts = 0


class OcrStageError(DemoError):
    def __init__(
        self,
        code: str,
        message: str,
        ocr_results: dict[str, list[dict[str, Any]]],
        *,
        total_images: int,
        completed_images: int,
        duration_ms: int,
    ):
        super().__init__(code, message)
        self.ocr_results = ocr_results
        self.total_images = total_images
        self.completed_images = completed_images
        self.duration_ms = duration_ms


def ocr_worker_count(config: dict[str, Any]) -> int:
    value = config.get("ocr_workers", OCR_WORKERS)
    try:
        workers = int(value)
    except (TypeError, ValueError) as exc:
        raise DemoError("INVALID_OCR_WORKERS", "OCR 并发必须是整数") from exc
    if not 1 <= workers <= OCR_MAX_WORKERS:
        raise DemoError(
            "INVALID_OCR_WORKERS",
            f"OCR 并发必须在 1 到 {OCR_MAX_WORKERS} 之间",
        )
    return workers


def utc_now() -> datetime:
    return datetime.now(UTC)


def iso_now() -> str:
    return utc_now().isoformat(timespec="seconds")


def progress(message: str) -> None:
    print(f"[{datetime.now().astimezone().strftime('%H:%M:%S')}] {message}", flush=True)


def json_default(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f"Unsupported JSON type: {type(value).__name__}")


def dump_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )
    temporary.replace(path)


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def stable_hash(value: Any) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=json_default,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def natural_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def canonical_id(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    if not text or len(text) > 64 or not re.fullmatch(r"[A-Za-z0-9_-]+", text):
        raise DemoError("INVALID_PRODUCT_ID", f"非法商品ID: {text!r}")
    return text


def as_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def positive_decimal(value: Any) -> Decimal | None:
    result = as_decimal(value)
    return result if result is not None and result > 0 else None


def round_money(value: Decimal | None, places: int = 4) -> Decimal | None:
    if value is None:
        return None
    quantum = Decimal(1).scaleb(-places)
    return value.quantize(quantum, rounding=ROUND_HALF_UP)


def decimal_text(value: Decimal) -> str:
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def first_non_empty(*values: Any) -> Any:
    for value in values:
        if value is not None and str(value).strip() not in {"", "nan", "None", "null"}:
            return value
    return None


class StateStore:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.artifacts = self.root / "artifacts"
        self.root.mkdir(parents=True, exist_ok=True)
        self.artifacts.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._attempt_id: str | None = None
        self._db = sqlite3.connect(self.root / "state.sqlite3", check_same_thread=False)
        self._db.row_factory = sqlite3.Row
        with self._db:
            self._db.execute("PRAGMA journal_mode=WAL")
            self._db.execute(
                """
                CREATE TABLE IF NOT EXISTS cache (
                    stage TEXT NOT NULL,
                    cache_key TEXT NOT NULL,
                    payload_json TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    expires_at TEXT,
                    PRIMARY KEY(stage, cache_key)
                )
                """
            )
            self._db.execute(
                """
                CREATE TABLE IF NOT EXISTS runs (
                    run_id TEXT PRIMARY KEY,
                    manifest_hash TEXT NOT NULL,
                    status TEXT NOT NULL,
                    output_dir TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )
            self._db.execute(
                """
                CREATE TABLE IF NOT EXISTS events (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    run_id TEXT NOT NULL,
                    attempt_id TEXT,
                    platform TEXT,
                    product_id TEXT,
                    stage TEXT NOT NULL,
                    status TEXT NOT NULL,
                    severity TEXT NOT NULL,
                    error_code TEXT,
                    message TEXT,
                    attempts INTEGER,
                    cached INTEGER NOT NULL DEFAULT 0,
                    duration_ms INTEGER,
                    source_table TEXT,
                    source_column TEXT,
                    source_time TEXT,
                    artifact_ref TEXT,
                    created_at TEXT NOT NULL
                )
                """
            )
            event_columns = {
                str(row[1]) for row in self._db.execute("PRAGMA table_info(events)").fetchall()
            }
            if "attempt_id" not in event_columns:
                self._db.execute("ALTER TABLE events ADD COLUMN attempt_id TEXT")

    def close(self) -> None:
        with self._lock:
            self._db.close()

    def cleanup_expired(self, days: int = RETENTION_DAYS) -> None:
        cutoff = utc_now() - timedelta(days=days)
        cutoff_text = cutoff.isoformat(timespec="seconds")
        with self._lock, self._db:
            expired = self._db.execute(
                """
                SELECT stage, cache_key, payload_json FROM cache
                WHERE (expires_at IS NOT NULL AND expires_at < ?) OR created_at < ?
                """,
                (iso_now(), cutoff_text),
            ).fetchall()
            self._db.execute(
                """
                DELETE FROM cache
                WHERE (expires_at IS NOT NULL AND expires_at < ?) OR created_at < ?
                """,
                (iso_now(), cutoff_text),
            )
        for row in expired:
            try:
                payload = json.loads(row["payload_json"])
                artifact = payload.get("artifact_ref") if isinstance(payload, dict) else None
                if artifact:
                    target = Path(artifact).resolve()
                    if target.is_relative_to(self.artifacts) and target.exists():
                        target.unlink()
            except (OSError, ValueError, TypeError):
                pass
        for child in self.artifacts.glob("**/*"):
            try:
                if child.is_file() and datetime.fromtimestamp(child.stat().st_mtime, UTC) < cutoff:
                    child.unlink()
            except OSError:
                pass

    def cache_get(self, stage: str, key: str) -> dict[str, Any] | None:
        with self._lock:
            row = self._db.execute(
                "SELECT payload_json, expires_at FROM cache WHERE stage=? AND cache_key=?",
                (stage, key),
            ).fetchone()
        if not row:
            return None
        if row["expires_at"] and row["expires_at"] < iso_now():
            return None
        payload = json.loads(row["payload_json"])
        if isinstance(payload, dict):
            refs = [value for key, value in payload.items() if key.endswith("_ref") and value]
            if any(Path(value).is_absolute() and not Path(value).exists() for value in refs):
                return None
        return payload

    def cache_put(
        self,
        stage: str,
        key: str,
        payload: dict[str, Any],
        ttl_hours: int | None = None,
    ) -> None:
        expires = (
            (utc_now() + timedelta(hours=ttl_hours)).isoformat(timespec="seconds")
            if ttl_hours
            else None
        )
        encoded = json.dumps(payload, ensure_ascii=False, default=json_default)
        with self._lock, self._db:
            self._db.execute(
                """
                INSERT INTO cache(stage, cache_key, payload_json, created_at, expires_at)
                VALUES(?,?,?,?,?)
                ON CONFLICT(stage, cache_key) DO UPDATE SET
                    payload_json=excluded.payload_json,
                    created_at=excluded.created_at,
                    expires_at=excluded.expires_at
                """,
                (stage, key, encoded, iso_now(), expires),
            )

    def write_artifact(self, stage: str, key: str, payload: Any) -> Path:
        safe_stage = re.sub(r"[^A-Za-z0-9_-]", "_", stage)
        target = self.artifacts / safe_stage / f"{key}.json"
        dump_json(target, payload)
        return target

    def write_text_artifact(self, stage: str, key: str, text: str, suffix: str = ".md") -> Path:
        safe_stage = re.sub(r"[^A-Za-z0-9_-]", "_", stage)
        target = self.artifacts / safe_stage / f"{key}{suffix}"
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_suffix(target.suffix + ".tmp")
        temporary.write_text(text, encoding="utf-8")
        temporary.replace(target)
        return target

    def choose_run(
        self, manifest_hash: str, output_root: Path, force_new: bool = False
    ) -> tuple[str, Path, bool]:
        row = None
        if not force_new:
            with self._lock:
                row = self._db.execute(
                    """
                    SELECT run_id, output_dir FROM runs
                    WHERE manifest_hash=? AND status!='complete'
                    ORDER BY created_at DESC LIMIT 1
                    """,
                    (manifest_hash,),
                ).fetchone()
        if row:
            return row["run_id"], Path(row["output_dir"]), True
        run_id = f"{utc_now().strftime('%Y%m%d-%H%M%S-%f')[:-3]}-{manifest_hash[:8]}"
        output_dir = (output_root / run_id).resolve()
        output_dir.mkdir(parents=True, exist_ok=True)
        now = iso_now()
        with self._lock, self._db:
            self._db.execute(
                "INSERT INTO runs VALUES(?,?,?,?,?,?)",
                (run_id, manifest_hash, "running", str(output_dir), now, now),
            )
        return run_id, output_dir, False

    def finish_run(self, run_id: str, status: str) -> None:
        with self._lock, self._db:
            self._db.execute(
                "UPDATE runs SET status=?, updated_at=? WHERE run_id=?",
                (status, iso_now(), run_id),
            )

    def begin_attempt(self) -> str:
        self._attempt_id = f"attempt-{utc_now().strftime('%Y%m%d-%H%M%S-%f')}"
        return self._attempt_id

    def event(
        self,
        run_id: str,
        stage: str,
        status: str,
        *,
        platform: str | None = None,
        product_id: str | None = None,
        severity: str = "info",
        error_code: str | None = None,
        message: str | None = None,
        attempts: int | None = None,
        cached: bool = False,
        duration_ms: int | None = None,
        source_table: str | None = None,
        source_column: str | None = None,
        source_time: str | None = None,
        artifact_ref: str | None = None,
    ) -> None:
        values = (
            run_id,
            self._attempt_id,
            platform,
            product_id,
            stage,
            status,
            severity,
            error_code,
            message,
            attempts,
            1 if cached else 0,
            duration_ms,
            source_table,
            source_column,
            source_time,
            artifact_ref,
            iso_now(),
        )
        with self._lock, self._db:
            self._db.execute(
                """
                INSERT INTO events(
                    run_id,attempt_id,platform,product_id,stage,status,severity,error_code,message,
                    attempts,cached,duration_ms,source_table,source_column,source_time,
                    artifact_ref,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                values,
            )

    def events(self, run_id: str, attempt_id: str | None = None) -> list[dict[str, Any]]:
        with self._lock:
            if attempt_id is None:
                rows = self._db.execute(
                    "SELECT * FROM events WHERE run_id=? ORDER BY id", (run_id,)
                ).fetchall()
            else:
                rows = self._db.execute(
                    "SELECT * FROM events WHERE run_id=? AND attempt_id=? ORDER BY id",
                    (run_id, attempt_id),
                ).fetchall()
        return [dict(row) for row in rows]


def validate_object(value: Any, field_specs: list[dict[str, Any]]) -> list[str]:
    if not isinstance(value, dict):
        return ["response must be a JSON object"]
    errors: list[str] = []
    expected_names = {spec["name"] for spec in field_specs}
    for extra in sorted(set(value) - expected_names):
        errors.append(f"unexpected field: {extra}")
    type_map = {
        "string": str,
        "number": (int, float),
        "boolean": bool,
        "array": list,
        "integer": int,
    }
    for spec in field_specs:
        name = spec["name"]
        if name not in value:
            errors.append(f"missing field: {name}")
            continue
        item = value[name]
        if item is None:
            if not spec.get("nullable", False):
                errors.append(f"{name} cannot be null")
            continue
        expected = spec["type"]
        expected_type = type_map[expected]
        if expected == "number" and isinstance(item, bool):
            errors.append(f"{name} must be number")
        elif not isinstance(item, expected_type):
            errors.append(f"{name} must be {expected}")
        elif expected == "array" and not all(isinstance(v, str) for v in item):
            errors.append(f"{name} array items must be string")
        elif expected == "number" and as_decimal(item) is None:
            errors.append(f"{name} must be finite")
        if spec.get("enum") and item not in spec["enum"]:
            errors.append(f"{name} must be one of {spec['enum']}")
        if isinstance(item, str) and spec.get("max_length") and len(item) > spec["max_length"]:
            errors.append(f"{name} exceeds max_length={spec['max_length']}")
        if expected == "number" and item is not None and spec.get("minimum") is not None:
            number = as_decimal(item)
            if number is not None and number < as_decimal(spec["minimum"]):
                errors.append(f"{name} must be >= {spec['minimum']}")
    min_daily = as_decimal(value.get("min_ri_fu_liang"))
    max_daily = as_decimal(value.get("max_ri_fu_liang"))
    if min_daily is not None and max_daily is not None and min_daily > max_daily:
        errors.append("min_ri_fu_liang cannot exceed max_ri_fu_liang")
    return errors


SERVING_UNIT_PATTERN = r"胶囊|粒|片|条|袋|包|颗|支|贴|丸|锭|滴"
MEASURE_UNIT_PATTERN = r"kg|mg|ml|mL|g|L|千克|毫克|毫升|克|升"
PACKAGE_UNIT_PATTERN = r"礼盒装|礼盒|瓶|盒|罐|桶|件"
UNIT_EQUIVALENCE = (
    frozenset({"片", "粒", "胶囊"}),
    frozenset({"袋", "包"}),
)


def parse_spec(raw: Any) -> dict[str, Any]:
    text = "" if raw is None else str(raw).strip()
    if not text or text.lower() in {"nan", "none", "null"}:
        return {"main_spec": None, "package": None, "total": None, "unit": None, "status": "missing"}
    normalized = re.sub(r"\s+", "", text.replace("／", "/").replace("X", "×").replace("x", "×").replace("*", "×"))
    number = r"(\d+(?:\.\d+)?)"
    serving = rf"({SERVING_UNIT_PATTERN})"
    measure = rf"({MEASURE_UNIT_PATTERN})"
    package_unit = rf"({PACKAGE_UNIT_PATTERN})"
    total: Decimal | None = None
    main_spec = package = unit = None

    # 15ml*80支：计价单位是“支”，不是1200ml。
    measured_items = re.search(rf"^{number}{measure}(?:/[^×]+)?×{number}{serving}", normalized, re.I)
    if measured_items:
        measure_count, measure_unit, item_count, unit = measured_items.groups()
        total = Decimal(item_count)
        main_spec = f"{decimal_text(total)}{unit}"
        package = f"{decimal_text(Decimal(measure_count))}{measure_unit}/{unit}"

    # 100片*3礼盒装、900g*2罐、900g*2。
    if total is None:
        multiplied = re.search(
            rf"^{number}({SERVING_UNIT_PATTERN}|{MEASURE_UNIT_PATTERN})×{number}{package_unit}?",
            normalized,
            re.I,
        )
        if multiplied:
            item_count, unit, package_count, package_name = multiplied.groups()
            item_decimal = Decimal(item_count)
            package_decimal = Decimal(package_count)
            total = item_decimal * package_decimal
            main_spec = f"{decimal_text(item_decimal)}{unit}"
            package = f"{decimal_text(package_decimal)}{package_name or '件'}"

    # 2盒*30粒。
    if total is None:
        package_first = re.search(rf"^{number}{package_unit}×{number}{serving}", normalized, re.I)
        if package_first:
            package_count, package_name, item_count, unit = package_first.groups()
            total = Decimal(package_count) * Decimal(item_count)
            main_spec = f"{decimal_text(Decimal(item_count))}{unit}"
            package = f"{decimal_text(Decimal(package_count))}{package_name}"

    # 60粒/瓶、80支/礼盒装。
    if total is None:
        single_package = re.search(rf"^{number}{serving}/{package_unit}", normalized, re.I)
        if single_package:
            item_count, unit, package_name = single_package.groups()
            total = Decimal(item_count)
            main_spec = f"{decimal_text(total)}{unit}"
            package = f"1{package_name}"

    if total is None:
        fallback = re.search(rf"{number}({SERVING_UNIT_PATTERN}|{MEASURE_UNIT_PATTERN})", normalized, re.I)
        if fallback:
            item_count, unit = fallback.groups()
            total = Decimal(item_count)
            main_spec = f"{decimal_text(total)}{unit}"
    if total is None:
        return {"main_spec": text, "package": None, "total": None, "unit": None, "status": "unparsed"}
    total_value: int | float = int(total) if total == total.to_integral_value() else float(total)
    return {
        "main_spec": main_spec,
        "package": package,
        "total": total_value,
        "unit": unit,
        "status": "parsed",
    }


def normalized_unit(unit: str) -> str:
    for group in UNIT_EQUIVALENCE:
        if unit in group:
            return sorted(group)[0]
    return unit.casefold()


def calculate_costs(
    price: Any,
    total: Any,
    min_daily: Any,
    max_daily: Any,
    spec_unit: str | None,
    daily_text: Any,
) -> dict[str, Any]:
    price_d = positive_decimal(price)
    total_d = positive_decimal(total)
    min_d = positive_decimal(min_daily)
    max_d = positive_decimal(max_daily)
    if price_d is None or total_d is None:
        return {"unit_price": None, "min_daily_cost": None, "max_daily_cost": None, "status": "missing_input"}
    daily = "" if daily_text is None else str(daily_text)
    daily_units = re.findall(rf"{SERVING_UNIT_PATTERN}|{MEASURE_UNIT_PATTERN}", daily, re.I)
    if (min_d is not None or max_d is not None) and spec_unit and daily_units:
        normalized_spec = normalized_unit(spec_unit)
        normalized_daily = {normalized_unit(item) for item in daily_units}
        if normalized_spec not in normalized_daily:
            return {"unit_price": round_money(price_d / total_d), "min_daily_cost": None, "max_daily_cost": None, "status": "unit_mismatch"}
    return {
        "unit_price": round_money(price_d / total_d),
        "min_daily_cost": round_money(price_d * min_d / total_d) if min_d else None,
        "max_daily_cost": round_money(price_d * max_d / total_d) if max_d else None,
        "status": "ok",
    }


def safe_message(value: Any) -> str:
    text = str(value) or type(value).__name__
    text = re.sub(
        r"(?i)(authorization|api[_-]?key|token|password)(\s*[:=]\s*)[^\s,;]+",
        r"\1\2[REDACTED]",
        text,
    )
    return re.sub(r"(?i)Bearer\s+[A-Za-z0-9._~+/=-]+", "Bearer [REDACTED]", text)[:2000]


def price_month_and_freshness(
    source_time: Any, *, as_of: date | None = None
) -> tuple[str | None, str]:
    match = re.search(r"(?<!\d)(\d{4})-?(\d{2})(?!\d)", str(source_time or ""))
    if not match:
        return None, "未知（无价格月份）"
    year, month = map(int, match.groups())
    try:
        date(year, month, 1)
    except ValueError:
        return None, "未知（价格月份无效）"
    current = as_of or date.today()
    age_months = current.year * 12 + current.month - (year * 12 + month)
    month_text = f"{year:04d}-{month:02d}"
    if age_months < 0:
        return month_text, "异常（未来月份）"
    if age_months == 0:
        return month_text, "最新（本月）"
    if age_months == 1:
        return month_text, "较新（1个月前）"
    return month_text, f"待关注（{age_months}个月前）"


def identity_key(platform: str, product_id: str) -> str:
    return f"{platform}{PRODUCT_KEY_SEPARATOR}{product_id}"


def canonical_platform(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    if not text or len(text) > 32 or not re.fullmatch(r"[A-Za-z0-9._-]+", text):
        raise DemoError("INVALID_PLATFORM", f"非法平台标识: {text!r}")
    return text


class ConcurrencyMeter:
    def __init__(self) -> None:
        self.current = 0
        self.peak = 0
        self._lock = threading.Lock()

    def __enter__(self) -> "ConcurrencyMeter":
        with self._lock:
            self.current += 1
            self.peak = max(self.peak, self.current)
        return self

    def __exit__(self, *_: Any) -> None:
        with self._lock:
            self.current -= 1


def increment_metric(config: dict[str, Any], name: str) -> None:
    metrics = config.setdefault("metrics", {})
    lock = config.setdefault("metrics_lock", threading.Lock())
    with lock:
        metrics[name] = metrics.get(name, 0) + 1


def run_parallel_stages(
    database_call: Callable[[], Any], ocr_call: Callable[[], Any]
) -> tuple[Any, Any, dict[str, int]]:
    def timed(call: Callable[[], Any]) -> tuple[Any, int]:
        started = time.perf_counter()
        result = call()
        return result, round((time.perf_counter() - started) * 1000)

    parallel_started = time.perf_counter()
    with futures.ThreadPoolExecutor(max_workers=2, thread_name_prefix="preflight") as executor:
        database_future = executor.submit(timed, database_call)
        ocr_future = executor.submit(timed, ocr_call)
        database_result, database_ms = database_future.result()
        ocr_result, ocr_ms = ocr_future.result()
    return database_result, ocr_result, {
        "redshift_ms": database_ms,
        "ocr_ms": ocr_ms,
        "parallel_wall_ms": round((time.perf_counter() - parallel_started) * 1000),
    }


def percentile_ms(values: list[int], percentile: int) -> int | None:
    if not values:
        return None
    ordered = sorted(values)
    index = max(0, (len(ordered) * percentile + 99) // 100 - 1)
    return ordered[index]


def build_performance_summary(
    stage_timings: dict[str, int],
    ocr_results: list[dict[str, Any]],
    metrics: dict[str, Any],
    peaks: dict[str, int],
    *,
    total_ms: int,
    planned_images: int | None = None,
) -> dict[str, Any]:
    total_images = len(ocr_results)
    success_images = sum(1 for result in ocr_results if result.get("ok"))
    cache_hits = sum(1 for result in ocr_results if result.get("cached"))
    durations = [
        int(result["duration_ms"])
        for result in ocr_results
        if not result.get("cached") and isinstance(result.get("duration_ms"), int) and result["duration_ms"] > 0
    ]
    success_durations = [
        int(result["duration_ms"])
        for result in ocr_results
        if result.get("ok")
        and not result.get("cached")
        and isinstance(result.get("duration_ms"), int)
        and result["duration_ms"] > 0
    ]
    image_inputs = [
        result["input_image"]
        for result in ocr_results
        if isinstance(result.get("input_image"), dict)
    ]
    source_bytes = sum(int(item.get("source_bytes", 0)) for item in image_inputs)
    sent_bytes = sum(int(item.get("sent_bytes", 0)) for item in image_inputs)
    resized_images = sum(
        1
        for item in image_inputs
        if (item.get("source_width"), item.get("source_height"))
        != (item.get("sent_width"), item.get("sent_height"))
    )
    ocr_ms = stage_timings.get("ocr_ms", 0)
    return {
        "total_ms": total_ms,
        "stages_ms": dict(stage_timings),
        "ocr": {
            "total_images": total_images,
            "planned_images": planned_images if planned_images is not None else total_images,
            "completed_images": total_images,
            "unstarted_images": max(0, (planned_images or total_images) - total_images),
            "success_images": success_images,
            "failed_images": total_images - success_images,
            "cache_hits": cache_hits,
            "api_calls": int(metrics.get("ocr_api_calls", 0)),
            "p50_ms": percentile_ms(durations, 50),
            "p95_ms": percentile_ms(durations, 95),
            "success_p50_ms": percentile_ms(success_durations, 50),
            "success_p95_ms": percentile_ms(success_durations, 95),
            "success_images_per_minute": round(success_images * 60_000 / ocr_ms, 2) if ocr_ms else None,
            "input_source_bytes": source_bytes,
            "input_sent_bytes": sent_bytes,
            "input_bytes_saved": source_bytes - sent_bytes,
            "resized_images": resized_images,
        },
        "metrics": dict(metrics),
        "peaks": dict(peaks),
    }


def compare_performance(current: dict[str, Any], previous: dict[str, Any]) -> dict[str, Any]:
    current_total = int(current.get("total_ms", 0))
    previous_total = int(previous.get("total_ms", 0))
    current_rate = current.get("ocr", {}).get("success_images_per_minute")
    previous_rate = previous.get("ocr", {}).get("success_images_per_minute")
    total_delta = current_total - previous_total
    rate_delta = None
    if isinstance(current_rate, (int, float)) and isinstance(previous_rate, (int, float)):
        rate_delta = round(float(current_rate) - float(previous_rate), 2)
    return {
        "previous_run_id": previous.get("run_id"),
        "total_ms_delta": total_delta,
        "total_ms_change_percent": round(total_delta * 100 / previous_total, 2) if previous_total else None,
        "ocr_success_images_per_minute_delta": rate_delta,
    }


def previous_performance(output_root: Path, current_output_dir: Path, run_fingerprint: str) -> dict[str, Any] | None:
    candidates: list[Path] = []
    paths = output_root.iterdir() if output_root.is_dir() else []
    for path in paths:
        if path == current_output_dir or not path.is_dir():
            continue
        performance_path = path / "performance.json"
        try:
            value = load_json(performance_path)
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(value, dict) and value.get("run_fingerprint") == run_fingerprint:
            candidates.append(performance_path)
    if not candidates:
        return None
    latest = max(candidates, key=lambda path: path.stat().st_mtime)
    value = load_json(latest)
    return value if isinstance(value, dict) else None


def post_json(
    url: str,
    payload: dict[str, Any],
    headers: dict[str, str],
    timeout: int,
) -> dict[str, Any]:
    request = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json", **headers},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read(4096).decode("utf-8", "replace")
        retryable = exc.code == 429 or 500 <= exc.code < 600
        raise DemoError(
            f"HTTP_{exc.code}",
            f"HTTP {exc.code}: {safe_message(body)}",
            retryable=retryable,
        ) from exc
    except (urllib.error.URLError, TimeoutError, OSError, http.client.HTTPException) as exc:
        raise DemoError("NETWORK_ERROR", safe_message(exc), retryable=True) from exc
    try:
        value = json.loads(raw)
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise DemoError("INVALID_RESPONSE_JSON", "接口返回的不是合法 JSON", retryable=True) from exc
    if not isinstance(value, dict):
        raise DemoError("INVALID_RESPONSE_SHAPE", "接口返回 JSON 不是对象", retryable=True)
    return value


def call_with_retry(
    call: Callable[[], Any],
    max_attempts: int,
    retry_delays: tuple[float, ...],
    on_attempt: Callable[[int], None] | None = None,
    on_retry: Callable[[int, DemoError, float], None] | None = None,
) -> tuple[Any, int]:
    if max_attempts < 1:
        raise ValueError("max_attempts must be at least 1")
    last_error: DemoError | None = None
    for attempt in range(1, max_attempts + 1):
        if on_attempt:
            on_attempt(attempt)
        try:
            return call(), attempt
        except DemoError as exc:
            last_error = exc
            exc.attempts = attempt
            if not exc.retryable or attempt == max_attempts:
                raise
            base_delay = retry_delays[min(attempt - 1, len(retry_delays) - 1)] if retry_delays else 0.0
            delay = base_delay + (random.uniform(0, min(1.0, base_delay * 0.1)) if base_delay else 0.0)
            if on_retry:
                on_retry(attempt, exc, delay)
            time.sleep(delay)
    raise last_error or DemoError("UNKNOWN", "未知重试错误")


def find_images(folder: Path) -> list[Path]:
    return sorted(
        [path for path in folder.rglob("*") if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS],
        key=lambda path: natural_key(path.relative_to(folder).as_posix()),
    )


def product_candidates(root: Path) -> list[Path]:
    return sorted(
        [path for path in root.iterdir() if path.is_dir() and find_images(path)],
        key=lambda path: natural_key(path.name),
    )


def build_manifest(
    root: Path, selected: list[Path], platform: str, *, allow_many: bool = False
) -> dict[str, Any]:
    root = root.resolve()
    platform = canonical_platform(platform)
    if not selected or (not allow_many and len(selected) > MAX_PRODUCTS):
        raise DemoError(
            "PRODUCT_COUNT",
            f"必须选择 1 到 {MAX_PRODUCTS} 个商品目录，当前为 {len(selected)} 个",
        )
    products: list[dict[str, Any]] = []
    seen: set[str] = set()
    for directory in selected:
        directory = directory.resolve()
        product_id = canonical_id(directory.name)
        if product_id in seen:
            raise DemoError("DUPLICATE_PRODUCT", f"商品 ID 重复: {product_id}")
        seen.add(product_id)
        images = find_images(directory)
        if not images:
            raise DemoError("NO_IMAGES", f"商品 {product_id} 目录中没有支持的图片")
        products.append(
            {
                "platform": platform,
                "product_id": product_id,
                "directory": str(directory),
                "images": [
                    {
                        "name": path.relative_to(directory).as_posix(),
                        "path": str(path.resolve()),
                        "sha256": file_sha256(path),
                        "bytes": path.stat().st_size,
                    }
                    for path in images
                ],
            }
        )
    products.sort(key=lambda item: natural_key(item["product_id"]))
    manifest = {"root": str(root), "platform": platform, "products": products}
    manifest["manifest_hash"] = stable_hash(
        [
            {
                "platform": product["platform"],
                "product_id": product["product_id"],
                "images": [(image["name"], image["sha256"]) for image in product["images"]],
            }
            for product in products
        ]
    )
    return manifest


def choose_root() -> Path:
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        selected = filedialog.askdirectory(title="选择包含 product_id 子目录的批次目录")
        root.destroy()
    except Exception as exc:
        raise DemoError("ROOT_REQUIRED", "无法打开目录选择器，请使用 --root 指定目录") from exc
    if not selected:
        raise DemoError("CANCELLED", "未选择批次目录")
    return Path(selected)


def choose_products(root: Path, requested: list[str] | None) -> list[Path]:
    candidates = product_candidates(root)
    if requested:
        lookup = {path.name: path for path in candidates}
        missing = [name for name in requested if name not in lookup]
        if missing:
            raise DemoError("PRODUCT_NOT_FOUND", f"目录不存在或没有图片: {', '.join(missing)}")
        return [lookup[name] for name in requested]
    print("\n可选商品目录：")
    for index, path in enumerate(candidates, 1):
        print(f"  {index:>2}. {path.name} ({len(find_images(path))} 张图片)")
    raw = input(f"请输入 1 到 {MAX_PRODUCTS} 个序号，用英文逗号分隔；输入 all 选择前 {MAX_PRODUCTS} 个: ").strip()
    try:
        indexes = (
            list(range(1, min(len(candidates), MAX_PRODUCTS) + 1))
            if raw.casefold() == "all"
            else [int(value.strip()) for value in raw.split(",") if value.strip()]
        )
        selected = [candidates[index - 1] for index in indexes]
    except (ValueError, IndexError) as exc:
        raise DemoError("INVALID_SELECTION", "商品目录序号无效") from exc
    if not 1 <= len(indexes) <= MAX_PRODUCTS or len(set(indexes)) != len(indexes):
        raise DemoError("PRODUCT_COUNT", f"必须选择 1 到 {MAX_PRODUCTS} 个不同的商品目录")
    return selected


def vllm_ocr_cache_key(image_sha256: str, model: str, model_version: str) -> str:
    return stable_hash(
        {
            "image_sha256": image_sha256,
            "provider": "vllm-qwen-vl",
            "model": model,
            "model_version": model_version,
            "prompt_version": VLLM_OCR_PROMPT_VERSION,
        }
    )


def vllm_markdown(response: dict[str, Any]) -> str:
    try:
        content = response["choices"][0]["message"]["content"]
        value = json.loads(content or "")
        blocks = value["blocks"]
    except (KeyError, IndexError, TypeError, json.JSONDecodeError) as exc:
        raise DemoError("INVALID_VLLM_OCR_RESPONSE", "本地 Qwen-VL 未返回可解析的 OCR JSON", retryable=True) from exc
    if not isinstance(blocks, list) or not blocks:
        raise DemoError("EMPTY_VLLM_OCR_BLOCKS", "本地 Qwen-VL 未返回文本块", retryable=True)
    ordered: list[tuple[int, int, str]] = []
    for index, block in enumerate(blocks):
        if not isinstance(block, dict):
            raise DemoError("INVALID_VLLM_OCR_BLOCK", f"第 {index + 1} 个文本块格式无效", retryable=True)
        order = block.get("order")
        text = block.get("text")
        if not isinstance(order, int) or order < 1 or not isinstance(text, str) or not text.strip():
            raise DemoError("INVALID_VLLM_OCR_BLOCK", f"第 {index + 1} 个文本块缺少 order 或 text", retryable=True)
        ordered.append((order, index, text.strip()))
    return "\n".join(text for _, _, text in sorted(ordered))


def vllm_log_id(response: dict[str, Any]) -> str | None:
    value = response.get("id")
    return str(value) if value is not None else None


def post_json_without_proxy(
    url: str,
    payload: dict[str, Any],
    headers: dict[str, str],
    timeout: int,
) -> dict[str, Any]:
    request = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json", **headers},
        method="POST",
    )
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(request, timeout=timeout) as response:
            raw = response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read(4096).decode("utf-8", "replace")
        retryable = exc.code == 429 or 500 <= exc.code < 600
        raise DemoError(f"VLLM_HTTP_{exc.code}", f"HTTP {exc.code}: {safe_message(body)}", retryable=retryable) from exc
    except (urllib.error.URLError, TimeoutError, OSError, http.client.HTTPException) as exc:
        raise DemoError(
            "VLLM_NETWORK_ERROR",
            "无法连接本地 Qwen-VL；请关闭市北 VPN 后重试，并确认在内网。" + safe_message(exc),
            retryable=True,
        ) from exc
    try:
        value = json.loads(raw)
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise DemoError("INVALID_VLLM_RESPONSE_JSON", "本地 Qwen-VL 返回的不是合法 JSON", retryable=True) from exc
    if not isinstance(value, dict):
        raise DemoError("INVALID_VLLM_RESPONSE_SHAPE", "本地 Qwen-VL 返回 JSON 不是对象", retryable=True)
    return value


def verify_vllm_available(api_base: str, api_key: str) -> None:
    request = urllib.request.Request(
        api_base.rstrip("/") + "/models",
        headers={"Authorization": f"Bearer {api_key}"},
        method="GET",
    )
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(request, timeout=10) as response:
            response.read()
    except (
        urllib.error.HTTPError,
        urllib.error.URLError,
        TimeoutError,
        OSError,
        http.client.HTTPException,
    ) as exc:
        raise DemoError(
            "VLLM_UNAVAILABLE",
            "无法连接本地 Qwen-VL；请关闭市北 VPN 后重试，并确认在内网。" + safe_message(exc),
        ) from exc


def prepare_vllm_image(image_path: Path) -> dict[str, Any]:
    source_bytes = image_path.stat().st_size
    with Image.open(image_path) as source:
        oriented = ImageOps.exif_transpose(source)
        source_width, source_height = oriented.size
        scale = min(
            1.0,
            OCR_IMAGE_MAX_LONG_EDGE / max(source_width, source_height),
            (OCR_IMAGE_MAX_PIXELS / (source_width * source_height)) ** 0.5,
        )
        sent_width = max(1, int(source_width * scale))
        sent_height = max(1, int(source_height * scale))
        image = oriented.resize((sent_width, sent_height), Image.Resampling.LANCZOS) if scale < 1 else oriented.copy()
        if "A" in image.getbands():
            background = Image.new("RGBA", image.size, "white")
            background.alpha_composite(image.convert("RGBA"))
            image = background.convert("RGB")
        elif image.mode != "RGB":
            image = image.convert("RGB")
        encoded = io.BytesIO()
        image.save(encoded, format="JPEG", quality=OCR_IMAGE_JPEG_QUALITY, optimize=True)
    image_bytes = encoded.getvalue()
    return {
        "bytes": image_bytes,
        "mime_type": "image/jpeg",
        "source_width": source_width,
        "source_height": source_height,
        "source_bytes": source_bytes,
        "sent_width": sent_width,
        "sent_height": sent_height,
        "sent_bytes": len(image_bytes),
    }


def vllm_ocr_payload_from_prepared_image(prepared: dict[str, Any], model: str) -> dict[str, Any]:
    image_url = (
        f"data:{prepared['mime_type']};base64,"
        f"{base64.b64encode(prepared['bytes']).decode('ascii')}"
    )
    return {
        "model": model,
        "messages": [
            {"role": "system", "content": VLLM_OCR_SYSTEM_PROMPT},
            {
                "role": "user",
                "content": [
                    {"type": "image_url", "image_url": {"url": image_url}},
                    {"type": "text", "text": "请转写这张商品图片中的全部可见文字。"},
                ],
            },
        ],
        "temperature": 0.0,
        "max_tokens": 6000,
        "response_format": {
            "type": "json_schema",
            "json_schema": {"name": "product_package_ocr", "strict": True, "schema": VLLM_OCR_SCHEMA},
        },
        "chat_template_kwargs": {"enable_thinking": False},
    }


def vllm_ocr_payload(image_path: Path, model: str) -> dict[str, Any]:
    return vllm_ocr_payload_from_prepared_image(prepare_vllm_image(image_path), model)


def paddle_request_id(product: dict[str, Any], image: dict[str, Any]) -> str:
    """单次 OCR 请求的唯一标识。

    这里刻意**不能**只用 ``image['sha256']``：同一个商品目录里完全可能有两张
    字节一致的图片（例如 20.jpg 和 26.jpg 是同一张素材），那样两张图会拿到
    同一个 request_id，进而引发两个后果：

    1. 同一个批次里出现重复 id，PaddleOCR 服务端直接以 HTTP 422 拒收整批，
       同批其余 7 张本来正常的图片一起被判为待复核；
    2. ``product_by_request`` 以 request_id 为键，后写入的那张图会覆盖前一张，
       导致前一张的图片名被替换掉（20.jpg 从结果里消失、26.jpg 出现两次）。

    图片的 ``name`` 在商品目录内唯一（相对路径 + 自然排序），因此用它来区分。
    """
    return f"{product['platform']}/{product['product_id']}/{image['name']}"


def paddle_ocr_cache_key(image_sha256: str, model_version: str) -> str:
    return stable_hash({
        "image_sha256": image_sha256,
        "provider": "paddleocr-vl",
        "model_version": model_version,
        "payload_version": PADDLE_OCR_PROMPT_VERSION,
    })


def paddle_result_to_ocr_result(
    image: dict[str, Any], result: dict[str, Any], *, duration_ms: int, attempts: int,
) -> dict[str, Any]:
    markdown = str(result.get("text") or "").strip()
    base = {
        "image_name": image["name"], "image_path": image["path"], "sha256": image["sha256"],
        "duration_ms": duration_ms, "attempts": attempts, "cached": False,
        "log_id": str(result.get("id") or "") or None, "input_image": None,
    }
    if not markdown:
        return {**base, "ok": False, "markdown": None, "raw_json_ref": None, "markdown_ref": None,
                "error": {"code": "PADDLE_OCR_REVIEW", "message": "PaddleOCR 未返回可用文字"}}
    return {**base, "ok": True, "markdown": markdown, "raw_json_ref": None, "markdown_ref": None, "error": None}


def paddle_worker_count(config: dict[str, Any]) -> int:
    try:
        workers = int(config.get("paddle_workers", 1))
    except (TypeError, ValueError) as exc:
        raise DemoError("INVALID_PADDLE_WORKERS", "Paddle 并行数只能是 1 或 2") from exc
    if workers not in {1, 2}:
        raise DemoError("INVALID_PADDLE_WORKERS", "Paddle 并行数只能是 1 或 2")
    return workers


def is_paddle_service_interruption(error: PaddleOcrError) -> bool:
    message = safe_message(error).casefold()
    return "paddle ocr unavailable:" in message or any(
        f"paddle ocr http {status}" in message for status in (500, 502, 503, 504)
    )


def write_paddle_interruption_record(
    run_id: str,
    config: dict[str, Any],
    reason: str,
    attempts: int,
    delay_seconds: float,
) -> str | None:
    output_text = str(config.get("run_output_dir") or "").strip()
    if not output_text:
        return None
    try:
        output_dir = Path(output_text)
        output_dir.mkdir(parents=True, exist_ok=True)
        target = output_dir / f"ocr服务中断-{utc_now().strftime('%Y%m%d-%H%M%S-%f')}.json"
        dump_json(
            target,
            {
                "run_id": run_id,
                "detected_at": iso_now(),
                "reason": reason,
                "automatic_recovery_attempts": attempts,
                "automatic_recovery_delay_seconds": delay_seconds,
                "manual_resume": "服务恢复后重新测试 OCR 服务，再点击重新运行；不要勾选强制创建新运行。",
            },
        )
        return str(target)
    except OSError:
        return None


def wait_for_paddle_recovery(
    run_id: str,
    config: dict[str, Any],
    store: StateStore,
    error: PaddleOcrError,
) -> bool:
    try:
        attempts = int(config.get("paddle_recovery_attempts", PADDLE_RECOVERY_ATTEMPTS))
        delay_seconds = float(
            config.get("paddle_recovery_delay_seconds", PADDLE_RECOVERY_DELAY_SECONDS)
        )
    except (TypeError, ValueError) as exc:
        raise DemoError("INVALID_PADDLE_RECOVERY", "Paddle 服务恢复参数无效") from exc
    if attempts < 1 or attempts > 60 or delay_seconds < 0 or delay_seconds > 60:
        raise DemoError("INVALID_PADDLE_RECOVERY", "Paddle 服务恢复参数超出允许范围")

    reason = safe_message(error)
    record_ref = write_paddle_interruption_record(
        run_id, config, reason, attempts, delay_seconds
    )
    store.event(
        run_id,
        "ocr_service",
        "interrupted",
        severity="error",
        error_code="PADDLE_OCR_INTERRUPTED",
        message=f"Paddle OCR 服务中断：{reason}",
        attempts=attempts,
        artifact_ref=record_ref,
    )
    progress(
        f"Paddle OCR 服务中断：{reason}；自动等待恢复（最多约 "
        f"{round((attempts - 1) * delay_seconds / 60, 1)} 分钟），成功图片已保存"
    )
    for attempt in range(1, attempts + 1):
        if attempt > 1:
            time.sleep(delay_seconds)
        try:
            verify_paddle_available(str(config["paddle_ocr_api_url"]))
        except PaddleOcrError:
            continue
        store.event(
            run_id,
            "ocr_service",
            "recovered",
            message=f"Paddle OCR 服务已恢复，第 {attempt} 次探测成功",
            attempts=attempt,
        )
        progress(f"Paddle OCR 服务已恢复，第 {attempt} 次探测成功；继续未完成图片")
        return True
    store.event(
        run_id,
        "ocr_service",
        "manual_resume_required",
        severity="error",
        error_code="PADDLE_OCR_INTERRUPTED",
        message="Paddle OCR 服务在自动恢复窗口内未恢复；本次运行已保留，需手动恢复",
        attempts=attempts,
    )
    return False


def raise_paddle_manual_resume_required() -> None:
    raise DemoError(
        "PADDLE_OCR_INTERRUPTED",
        "Paddle OCR 服务未在自动恢复窗口内恢复；成功图片已保存。服务恢复后，"
        "重新测试 OCR 服务，再点击“重新运行”（不要勾选强制创建新运行）即可继续。",
    )


def run_paddle_ocr_stage(
    run_id: str,
    manifest: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
    meter: ConcurrencyMeter,
    on_product_partial: Callable[[dict[str, Any], list[dict[str, Any]], list[str]], None] | None = None,
) -> dict[str, list[dict[str, Any]]]:
    output = {identity_key(product["platform"], product["product_id"]): [] for product in manifest["products"]}
    product_by_request: dict[str, tuple[dict[str, Any], dict[str, Any], str]] = {}
    pending: list[PaddleImage] = []
    model_version = str(config.get("paddle_ocr_model_version", PADDLE_DEFAULT_MODEL_VERSION))
    for product in manifest["products"]:
        key = identity_key(product["platform"], product["product_id"])
        for image in product["images"]:
            request_id = paddle_request_id(product, image)
            cache_key = paddle_ocr_cache_key(image["sha256"], model_version)
            product_by_request[request_id] = (product, image, cache_key)
            cached = None if config.get("force_ocr") else store.cache_get("ocr", cache_key)
            if cached:
                output[key].append({**cached, "image_name": image["name"], "image_path": image["path"], "sha256": image["sha256"], "cached": True})
            else:
                pending.append(PaddleImage(request_id, Path(image["path"])))
    batches = [pending[index:index + PADDLE_BATCH_SIZE] for index in range(0, len(pending), PADDLE_BATCH_SIZE)]
    total = sum(len(product["images"]) for product in manifest["products"])
    progress(f"Paddle OCR 阶段开始：共 {total} 张图片，批次 {len(batches)}，并行 {paddle_worker_count(config)}")
    if pending:
        try:
            verify_paddle_available(str(config["paddle_ocr_api_url"]))
        except PaddleOcrError as exc:
            if not wait_for_paddle_recovery(run_id, config, store, exc):
                raise_paddle_manual_resume_required()

    def post(url: str, images: list[PaddleImage]) -> list[dict[str, Any]]:
        increment_metric(config, "ocr_api_calls")
        with meter:
            return post_batch(url, images, timeout=30)

    def run_batch(
        batch: list[PaddleImage],
    ) -> tuple[list[PaddleImage], list[dict[str, Any]], int, int, dict[str, str]]:
        """请求一个批次。

        返回 ``(批次, 成功响应, 耗时ms, 尝试次数, 逐图错误映射)``。
        只有整批成功时错误映射才为空；批次失败时会退化成逐图请求，
        所以映射里只剩真正失败的那几张。
        """
        api_url = str(config["paddle_ocr_api_url"])
        started = time.perf_counter()
        error: str | None = None
        attempts = 0
        while True:
            last_error: PaddleOcrError | None = None
            for _ in range(2):
                attempts += 1
                try:
                    results = post(api_url, batch)
                    return batch, results, round((time.perf_counter() - started) * 1000), attempts, {}
                except PaddleOcrError as exc:
                    last_error = exc
                    error = safe_message(exc)
            if last_error is None or not is_paddle_service_interruption(last_error):
                break
            if not wait_for_paddle_recovery(run_id, config, store, last_error):
                raise_paddle_manual_resume_required()

        # 整批被拒时不再把同批所有图片一起判死：逐张重试，能救几张救几张。
        # 之前的行为是整批共用一个错误串，一张畸形图会带走另外 7 张正常图片。
        errors = {image.request_id: error or "Paddle OCR 请求失败" for image in batch}
        recovered: list[dict[str, Any]] = []
        if len(batch) > 1 and (last_error is None or not is_paddle_service_interruption(last_error)):
            progress(f"Paddle OCR 批次失败（{error}），改为逐张重试 {len(batch)} 张图片")
            for image in batch:
                try:
                    single = post(api_url, [image])
                except PaddleOcrError as exc:
                    errors[image.request_id] = safe_message(exc)
                    continue
                recovered.extend(single)
                errors.pop(image.request_id, None)
        return batch, recovered, round((time.perf_counter() - started) * 1000), attempts, errors

    completed = sum(len(items) for items in output.values())
    with futures.ThreadPoolExecutor(max_workers=paddle_worker_count(config), thread_name_prefix="paddle") as executor:
        for batch, results, duration_ms, attempts, error_map in executor.map(run_batch, batches):
            result_by_id = {str(item.get("id")): item for item in results}
            for requested in batch:
                product, image, cache_key = product_by_request[requested.request_id]
                key = identity_key(product["platform"], product["product_id"])
                if requested.request_id in error_map:
                    item = paddle_result_to_ocr_result(image, {}, duration_ms=duration_ms, attempts=attempts)
                    item["error"] = {
                        "code": "PADDLE_OCR_REVIEW",
                        "message": error_map[requested.request_id],
                    }
                else:
                    item = paddle_result_to_ocr_result(image, result_by_id[requested.request_id], duration_ms=duration_ms, attempts=attempts)
                raw_ref = store.write_artifact("ocr_raw" if item["ok"] else "ocr_failure", cache_key, result_by_id.get(requested.request_id, item))
                item["raw_json_ref"] = str(raw_ref)
                if item["ok"]:
                    item["markdown_ref"] = str(store.write_text_artifact("ocr_markdown", cache_key, item["markdown"]))
                    store.cache_put("ocr", cache_key, item)
                output[key].append(item)
                completed += 1
                store.event(run_id, "ocr", "success" if item["ok"] else "failed", platform=product["platform"], product_id=product["product_id"], severity="error" if not item["ok"] else "info", error_code=item["error"]["code"] if item["error"] else None, message=f"{image['name']} Paddle OCR {'完成' if item['ok'] else '待复核'}", attempts=attempts, duration_ms=duration_ms, artifact_ref=str(raw_ref))
                progress(f"Paddle OCR 进度 {completed}/{total}: {product['product_id']}/{image['name']} {'成功' if item['ok'] else '待复核'}")
    for product in manifest["products"]:
        key = identity_key(product["platform"], product["product_id"])
        output[key].sort(key=lambda item: natural_key(item["image_name"]))
        if on_product_partial:
            success = [item for item in output[key] if item.get("ok")]
            failed = [item["image_name"] for item in output[key] if not item.get("ok")]
            if success:
                on_product_partial(product, success, failed)
    return output


def run_ocr_one(
    run_id: str,
    product: dict[str, Any],
    image: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
    meter: ConcurrencyMeter,
    on_first_retry: Callable[[], None] | None = None,
) -> dict[str, Any]:
    cache_key = vllm_ocr_cache_key(
        image["sha256"], config["vllm_ocr_model"], config["vllm_ocr_model_version"]
    )
    cached = None if config.get("force_ocr") else store.cache_get("ocr", cache_key)
    if cached:
        store.event(
            run_id,
            "ocr",
            "success",
            platform=product["platform"],
            product_id=product["product_id"],
            message=f"{image['name']} 命中缓存",
            attempts=0,
            cached=True,
            duration_ms=0,
            artifact_ref=cached.get("raw_json_ref"),
        )
        return {
            **cached,
            "image_name": image["name"],
            "image_path": image["path"],
            "sha256": image["sha256"],
            "cached": True,
        }

    started = time.perf_counter()
    attempts = 0
    prepared_image: dict[str, Any] | None = None
    try:
        def request() -> tuple[dict[str, Any], str]:
            nonlocal prepared_image
            increment_metric(config, "ocr_api_calls")
            if config.get("mock"):
                if "fail_ocr" in image["name"].casefold():
                    raise DemoError("MOCK_OCR_FAILURE", "模拟 OCR 最终失败")
                response = {
                    "id": f"mock-{image['sha256'][:16]}",
                    "model": config["vllm_ocr_model"],
                    "choices": [{"message": {"content": json.dumps({
                        "blocks": [
                            {"order": 1, "text": f"# {product['product_id']} 演示包装", "legibility": "clear"},
                            {"order": 2, "text": "规格：60粒/瓶", "legibility": "clear"},
                            {"order": 3, "text": "建议每日2粒", "legibility": "clear"},
                            {"order": 4, "text": "成分：维生素C 100mg", "legibility": "clear"},
                        ],
                    }, ensure_ascii=False)}}],
                }
            else:
                if prepared_image is None:
                    prepared_image = prepare_vllm_image(Path(image["path"]))
                time_limit = float(config.get("ocr_request_timeout", OCR_REQUEST_TIMEOUT))
                remaining = time_limit - (time.perf_counter() - started)
                if remaining <= 0:
                    raise DemoError(
                        "VLLM_NETWORK_ERROR",
                        "单图 OCR 总等待时间已耗尽",
                    )
                request_started = time.perf_counter()
                try:
                    response = post_json_without_proxy(
                        config["vllm_ocr_api_base"].rstrip("/") + "/chat/completions",
                        vllm_ocr_payload_from_prepared_image(prepared_image, config["vllm_ocr_model"]),
                        {"Authorization": f"Bearer {config['vllm_ocr_api_key']}"},
                        timeout=max(1, int(remaining)),
                    )
                except DemoError as exc:
                    if exc.code == "VLLM_NETWORK_ERROR":
                        fast_failure_limit = float(
                            config.get(
                                "ocr_fast_network_failure_seconds",
                                OCR_FAST_NETWORK_FAILURE_SECONDS,
                            )
                        )
                        if time.perf_counter() - request_started > fast_failure_limit:
                            exc.retryable = False
                    else:
                        exc.retryable = False
                    raise
            try:
                markdown = vllm_markdown(response)
            except DemoError as exc:
                exc.retryable = False
                raise
            if not markdown:
                raise DemoError("EMPTY_OCR_TEXT", "本地 Qwen-VL 未返回文字")
            return response, markdown

        def on_retry(attempt: int, exc: DemoError, delay: float) -> None:
            if attempt == 1 and on_first_retry:
                on_first_retry()
            progress(
                f"OCR 暂时失败 {product['product_id']}/{image['name']}：{exc.code}；"
                f"{delay:.1f} 秒后进行第 {attempt + 1} 次尝试"
            )

        with meter:
            max_attempts = int(config.get("ocr_max_attempts", OCR_MAX_ATTEMPTS))
            retry_delays = tuple(config.get("ocr_retry_delays", OCR_RETRY_DELAYS))
            (response, markdown), attempts = call_with_retry(
                request,
                max_attempts=max_attempts,
                retry_delays=retry_delays,
                on_attempt=lambda attempt: progress(
                    f"OCR 请求 {product['product_id']}/{image['name']}，尝试 {attempt}/{max_attempts}"
                ),
                on_retry=on_retry,
            )
        duration_ms = round((time.perf_counter() - started) * 1000)
        raw_ref = store.write_artifact("ocr_raw", cache_key, response)
        markdown_ref = store.write_text_artifact("ocr_markdown", cache_key, markdown)
        result = {
            "ok": True,
            "image_name": image["name"],
            "image_path": image["path"],
            "sha256": image["sha256"],
            "log_id": vllm_log_id(response),
            "markdown": markdown,
            "duration_ms": duration_ms,
            "attempts": attempts + int(config.get("ocr_attempt_offset", 0)),
            "raw_json_ref": str(raw_ref),
            "markdown_ref": str(markdown_ref),
            "input_image": (
                {key: value for key, value in prepared_image.items() if key != "bytes"}
                if prepared_image else None
            ),
            "error": None,
            "cached": False,
        }
        store.cache_put("ocr", cache_key, result)
        store.event(
            run_id,
            "ocr",
            "success",
            platform=product["platform"],
            product_id=product["product_id"],
            message=f"{image['name']} OCR 完成，logId={result['log_id'] or '-'}",
            attempts=attempts,
            duration_ms=duration_ms,
            artifact_ref=str(raw_ref),
        )
        return result
    except DemoError as exc:
        duration_ms = round((time.perf_counter() - started) * 1000)
        failure = {
            "ok": False,
            "image_name": image["name"],
            "image_path": image["path"],
            "sha256": image["sha256"],
            "log_id": None,
            "markdown": None,
            "duration_ms": duration_ms,
            "input_image": (
                {key: value for key, value in prepared_image.items() if key != "bytes"}
                if prepared_image else None
            ),
            "attempts": (exc.attempts or attempts or 1) + int(config.get("ocr_attempt_offset", 0)),
            "error": {"code": exc.code, "message": safe_message(exc)},
            "cached": False,
        }
        failure_ref = store.write_artifact(
            "ocr_failure", f"{cache_key}-{run_id}", failure
        )
        failure["raw_json_ref"] = str(failure_ref)
        failure["markdown_ref"] = None
        store.event(
            run_id,
            "ocr",
            "failed",
            platform=product["platform"],
            product_id=product["product_id"],
            severity="error",
            error_code=exc.code,
            message=f"{image['name']}: {safe_message(exc)}",
            attempts=failure["attempts"],
            duration_ms=duration_ms,
            artifact_ref=str(failure_ref),
        )
        return failure


def run_ocr_stage(
    run_id: str,
    manifest: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
    meter: ConcurrencyMeter,
    on_product_partial: Callable[[dict[str, Any], list[dict[str, Any]], list[str]], None] | None = None,
) -> dict[str, list[dict[str, Any]]]:
    output = {
        identity_key(product["platform"], product["product_id"]): []
        for product in manifest["products"]
    }
    jobs: dict[futures.Future[dict[str, Any]], tuple[dict[str, Any], dict[str, Any]]] = {}
    partial_lock = threading.Lock()
    partial_emitted: set[str] = set()
    ocr_workers = ocr_worker_count(config)
    active_workers = ocr_workers
    recovery_successes = 0
    total = sum(len(product["images"]) for product in manifest["products"])
    completed = 0
    stage_started = time.perf_counter()
    progress(f"OCR 阶段开始：共 {total} 张图片，最大并发 {ocr_workers}")

    pending_jobs = [
        (product, image)
        for product in manifest["products"]
        for image in product["images"]
    ]
    next_job = 0
    consecutive_network_failures = 0
    outage_detected = False
    completed_futures: queue.SimpleQueue[futures.Future[dict[str, Any]]] = queue.SimpleQueue()
    with futures.ThreadPoolExecutor(max_workers=ocr_workers, thread_name_prefix="ocr") as executor:
        def submit_next() -> bool:
            nonlocal next_job
            if next_job >= len(pending_jobs):
                return False
            product, image = pending_jobs[next_job]
            next_job += 1
            foreground_config = dict(config)
            foreground_config["ocr_max_attempts"] = 2
            foreground_config["ocr_retry_delays"] = (OCR_FAST_NETWORK_RETRY_DELAY,)
            foreground_config["ocr_request_timeout"] = int(
                config.get("ocr_request_timeout", OCR_REQUEST_TIMEOUT)
            )
            future = executor.submit(
                run_ocr_one,
                run_id,
                product,
                image,
                foreground_config,
                store,
                meter,
            )
            jobs[future] = (product, image)
            future.add_done_callback(completed_futures.put)
            return True

        def replace_image_result(key: str, image_name: str, result: dict[str, Any]) -> None:
            with partial_lock:
                for index, current in enumerate(output[key]):
                    if current.get("image_name") == image_name:
                        output[key][index] = result
                        return
                output[key].append(result)

        for _ in range(min(active_workers, len(pending_jobs))):
            submit_next()
        while jobs:
            future = completed_futures.get()
            product, image = jobs.pop(future)
            result = future.result()
            key = identity_key(product["platform"], product["product_id"])
            with partial_lock:
                output[key].append(result)
            completed += 1
            progress(
                f"OCR 进度 {completed}/{total}：{product['product_id']}/{image['name']} "
                f"{'成功' if result.get('ok') else '失败'}"
            )
            is_network_failure = (
                not result.get("ok")
                and result.get("error", {}).get("code") == "VLLM_NETWORK_ERROR"
            )
            if is_network_failure:
                consecutive_network_failures += 1
                active_workers = 1
                recovery_successes = 0
                progress("本地 OCR 出现网络失败，后续请求降为单并发以保护服务")
                outage_detected = (
                    consecutive_network_failures >= VLLM_NETWORK_FAILURE_THRESHOLD
                )
            else:
                consecutive_network_failures = 0
                if result.get("ok") and active_workers < min(2, ocr_workers):
                    recovery_successes += 1
                    if recovery_successes >= OCR_RECOVERY_SUCCESS_COUNT:
                        active_workers = min(2, ocr_workers)
                        recovery_successes = 0
                        progress("本地 OCR 已连续成功，恢复至双并发")
            if is_network_failure and not outage_detected:
                timeout_result = {
                    **result,
                    "error": {
                        "code": "OCR_TIMEOUT_REVIEW",
                        "message": "单图已超过 30 秒上限，已放弃重试并标记待复核",
                    },
                }
                replace_image_result(key, image["name"], timeout_result)
                progress(f"OCR 超时待复核：{product['product_id']}/{image['name']}；继续后续图片")
            product_total = len(product["images"])
            if len(output[key]) == product_total and on_product_partial:
                with partial_lock:
                    should_emit = key not in partial_emitted
                    successful_images = [item for item in output[key] if item.get("ok")]
                    failed_images = [item["image_name"] for item in output[key] if not item.get("ok")]
                    if should_emit and successful_images:
                        partial_emitted.add(key)
                if should_emit and successful_images:
                    on_product_partial(product, list(successful_images), failed_images)
            while not outage_detected and len(jobs) < active_workers and submit_next():
                pass
    if outage_detected:
        raise OcrStageError(
            "VLLM_NETWORK_OUTAGE",
            "本地 Qwen-VL 连续 3 张图片网络失败，已停止剩余图片；请检查市北 VPN、内网和服务端后重试。",
            output,
            total_images=total,
            completed_images=completed,
            duration_ms=round((time.perf_counter() - stage_started) * 1000),
        )
    for product in manifest["products"]:
        key = identity_key(product["platform"], product["product_id"])
        output[key].sort(key=lambda item: natural_key(item["image_name"]))
    return output


def extract_json_object(content: str) -> dict[str, Any]:
    text = content.strip()
    if "</think>" in text:
        text = text.split("</think>", 1)[1].strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
        text = re.sub(r"\s*```$", "", text)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < start:
        raise DemoError("INVALID_MODEL_JSON", "模型没有返回 JSON 对象", retryable=True)
    try:
        value = json.loads(
            text[start : end + 1],
            parse_constant=lambda constant: (_ for _ in ()).throw(ValueError(constant)),
        )
    except (json.JSONDecodeError, ValueError) as exc:
        raise DemoError("INVALID_MODEL_JSON", "模型返回 JSON 无法解析", retryable=True) from exc
    if not isinstance(value, dict):
        raise DemoError("INVALID_MODEL_JSON", "模型返回值不是 JSON 对象", retryable=True)
    return value


def provider_search_sources(response: dict[str, Any]) -> list[dict[str, str]]:
    """Read provider-returned search metadata; never trust URLs written in model content."""
    found: dict[str, dict[str, str]] = {}

    def collect(value: Any) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                if key.casefold() in {"search_results", "searchresults", "web_search_results"}:
                    for item in child if isinstance(child, list) else []:
                        if not isinstance(item, dict):
                            continue
                        url = first_non_empty(item.get("url"), item.get("link"))
                        parsed = urllib.parse.urlparse(str(url or ""))
                        if parsed.scheme in {"http", "https"} and parsed.netloc:
                            found[str(url)] = {
                                "url": str(url),
                                "title": str(first_non_empty(item.get("title"), item.get("name")) or ""),
                            }
                collect(child)
        elif isinstance(value, list):
            for child in value:
                collect(child)

    collect(response)
    return list(found.values())


def mock_qwen_values(field_specs: list[dict[str, Any]], search: bool) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for spec in field_specs:
        if spec["type"] == "boolean":
            values[spec["name"]] = False
        elif spec["type"] == "array":
            values[spec["name"]] = []
        elif spec.get("enum"):
            values[spec["name"]] = "neutral" if "neutral" in spec["enum"] else spec["enum"][0]
        else:
            values[spec["name"]] = None
    if search:
        values.update({"factory_name": "模拟生产方", "origin": "中国", "hot_topics": "本地演示"})
    else:
        values.update(
            {
                "guige": "60粒/瓶",
                "guige_zong_liang": 60,
                "ri_fu_liang": "每日2粒",
                "min_ri_fu_liang": 2,
                "max_ri_fu_liang": 2,
                "ingredients": "维生素C",
                "gender_tendency": "neutral",
                "brand": "演示品牌",
                "dosage_form": "片剂",
                "ingredient_content": "维生素C 100mg",
            }
        )
    return values


def run_qwen_one(
    run_id: str,
    product: dict[str, Any],
    kind: str,
    name: str,
    ocr_text: str,
    attributes: str,
    db_snapshot_hash: str,
    template: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
    stage_meter: ConcurrencyMeter,
    total_meter: ConcurrencyMeter,
) -> dict[str, Any]:
    is_search = kind == "search"
    section = template["search"] if is_search else template["model"]
    user_prompt = section["user_prompt_template"].format(
        name=name,
        ocr_text=ocr_text,
        attributes=attributes or "无登记属性",
    )
    cache_key = stable_hash(
        {
            "kind": kind,
            "model": section["model"] if is_search else section["name"],
            "system": section["system_prompt"],
            "user": user_prompt,
            "fields": section["fields"],
            "db_snapshot_hash": None if is_search else db_snapshot_hash,
            "response_parser": "qwen-search-evidence-v1" if is_search else "qwen-extract-v1",
        }
    )
    cached = store.cache_get(f"qwen_{kind}", cache_key)
    if cached:
        store.event(
            run_id,
            f"qwen_{kind}",
            "success",
            platform=product["platform"],
            product_id=product["product_id"],
            message="命中缓存",
            attempts=0,
            cached=True,
            duration_ms=0,
            artifact_ref=cached.get("raw_json_ref"),
        )
        return {**cached, "cached": True}

    started = time.perf_counter()
    attempts = 0
    last_raw_response: dict[str, Any] | None = None
    try:
        def request() -> tuple[dict[str, Any], dict[str, Any]]:
            nonlocal last_raw_response
            increment_metric(config, f"qwen_{kind}_api_calls")
            if config.get("mock"):
                parsed = mock_qwen_values(section["fields"], is_search)
                raw_response = {"choices": [{"message": {"content": json.dumps(parsed, ensure_ascii=False)}}]}
                if is_search:
                    raw_response["search_info"] = {
                        "search_results": [{"title": "模拟来源", "url": "https://example.com/mock-source"}]
                    }
            else:
                payload = {
                    "model": section["model"] if is_search else section["name"],
                    "messages": [
                        {"role": "system", "content": section["system_prompt"]},
                        {"role": "user", "content": user_prompt},
                    ],
                    "temperature": section.get("temperature", 0.3),
                    "max_tokens": section.get("max_tokens", 2000),
                }
                if is_search:
                    payload["enable_search"] = True
                raw_response = post_json(
                    config.get("qwen_base_url", QWEN_BASE_URL).rstrip("/") + "/chat/completions",
                    payload,
                    {"Authorization": f"Bearer {config['qwen_api_key']}"},
                    timeout=90,
                )
                last_raw_response = raw_response
                try:
                    content = raw_response["choices"][0]["message"]["content"]
                except (KeyError, IndexError, TypeError) as exc:
                    raise DemoError("INVALID_QWEN_RESPONSE", "Qwen 响应缺少 message.content", retryable=True) from exc
                parsed = extract_json_object(str(content))
            errors = validate_object(parsed, section["fields"])
            if errors:
                raise DemoError("INVALID_QWEN_FIELDS", "; ".join(errors), retryable=True)
            return parsed, raw_response

        with total_meter, stage_meter:
            max_attempts = int(config.get("qwen_max_attempts", QWEN_MAX_ATTEMPTS))
            retry_delays = tuple(config.get("qwen_retry_delays", QWEN_RETRY_DELAYS))
            (parsed, raw_response), attempts = call_with_retry(
                request,
                max_attempts=max_attempts,
                retry_delays=retry_delays,
                on_attempt=lambda attempt: progress(
                    f"Qwen {kind} 请求 {product['product_id']}，尝试 {attempt}/{max_attempts}"
                ),
                on_retry=lambda attempt, exc, delay: progress(
                    f"Qwen {kind} 暂时失败 {product['product_id']}：{exc.code}；"
                    f"{delay:.1f} 秒后进行第 {attempt + 1} 次尝试"
                ),
            )
        duration_ms = round((time.perf_counter() - started) * 1000)
        raw_ref = store.write_artifact(f"qwen_{kind}_raw", cache_key, raw_response)
        result = {
            "ok": True,
            "data": parsed,
            "provider_sources": provider_search_sources(raw_response) if is_search else [],
            "duration_ms": duration_ms,
            "attempts": attempts,
            "raw_json_ref": str(raw_ref),
            "error": None,
            "cached": False,
        }
        ttl = int(section.get("cache_ttl_hours", 24)) if is_search else None
        store.cache_put(f"qwen_{kind}", cache_key, result, ttl_hours=ttl)
        store.event(
            run_id,
            f"qwen_{kind}",
            "success",
            platform=product["platform"],
            product_id=product["product_id"],
            attempts=attempts,
            duration_ms=duration_ms,
            artifact_ref=str(raw_ref),
        )
        return result
    except DemoError as exc:
        duration_ms = round((time.perf_counter() - started) * 1000)
        severity = "warning" if is_search else "error"
        failure = {
            "ok": False,
            "kind": kind,
            "duration_ms": duration_ms,
            "attempts": exc.attempts or attempts or 1,
            "error": {"code": exc.code, "message": safe_message(exc)},
            "last_response": last_raw_response,
            "cached": False,
        }
        failure_ref = store.write_artifact(
            f"qwen_{kind}_failure", f"{cache_key}-{run_id}", failure
        )
        store.event(
            run_id,
            f"qwen_{kind}",
            "failed",
            platform=product["platform"],
            product_id=product["product_id"],
            severity=severity,
            error_code=exc.code,
            message=safe_message(exc),
            attempts=failure["attempts"],
            duration_ms=duration_ms,
            artifact_ref=str(failure_ref),
        )
        return {**failure, "data": {}, "raw_json_ref": str(failure_ref)}


def mock_database_results(manifest: dict[str, Any]) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    for index, product in enumerate(manifest["products"]):
        key_value = 22129910 if product["product_id"] == "497394" else 30000000 + index
        price = Decimal("99.90") + Decimal(index)
        source_label = (
            "京东（主平台，platform_key=1）"
            if product["platform"].casefold() == "jd"
            else f"{product['platform']}（模拟平台）"
        )
        record = {
            "identity": {
                "platform": product["platform"],
                "product_id": product["product_id"],
                "platform_match_mode": "exact",
                "candidate_row_count": 1,
                "scoped_row_count": 1,
                "snapshot_at": iso_now(),
            },
            "key": {
                "value": key_value,
                "status": "ok",
                "source_table": "d_platform_goods",
                "source_time": "2026-08-01T00:00:00+00:00",
                "goods_value": key_value,
                "monthly_value": key_value,
                "goods_variant_count": 1,
                "monthly_variant_count": 1,
                "source_conflict": False,
            },
            "name": {
                "value": f"演示商品 {product['product_id']} 60粒/瓶",
                "source_table": "d_platform_goods",
                "source_time": "2026-08-01T00:00:00+00:00",
            },
            "platform_source": {"key": 1, "label": source_label},
            "attributes": [
                {
                    "name": "品牌",
                    "value": "演示品牌",
                    "source_table": "d_platform_goods_attributes",
                    "source_time": "2026-08-01T00:00:00+00:00",
                },
                {
                    "name": "剂型",
                    "value": "片剂",
                    "source_table": "d_platform_goods_attributes",
                    "source_time": "2026-08-01T00:00:00+00:00",
                },
            ],
            "price": {
                "value": price,
                "raw_value": price,
                "status": "ok",
                "source_table": "mv_com_goods_statistics_monthly_v2_internal_ssv4",
                "source_column": "lowest_promo_price",
                "source_time": "2026-08-01T00:00:00+00:00",
                "platform_source": {"key": 1, "label": source_label},
                "monthly_variant_count": 1,
                "goods_variant_count": 0,
            },
            "review_issues": [],
        }
        record["db_snapshot_hash"] = stable_hash(record)
        output[identity_key(product["platform"], product["product_id"])] = record
    return output


def run_database_stage(
    run_id: str,
    manifest: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
) -> tuple[dict[str, dict[str, Any]], str]:
    progress(f"PostgreSQL 批量查询开始：{len(manifest['products'])} 个商品")
    identities = [
        {"platform": item["platform"], "product_id": item["product_id"]}
        for item in manifest["products"]
    ]
    cache_key = stable_hash(
        {"run_id": run_id, "identities": identities, "rule": SQL_RULE_VERSION}
    )
    cached = store.cache_get("redshift", cache_key)
    if cached:
        progress("PostgreSQL 命中当次数据库快照缓存")
        result = cached["result"]
        artifact_ref = cached["artifact_ref"]
        for product in manifest["products"]:
            key = identity_key(product["platform"], product["product_id"])
            record = result.get(key, {})
            price = record.get("price", {})
            store.event(
                run_id,
                "redshift",
                "success",
                platform=product["platform"],
                product_id=product["product_id"],
                message="恢复当次数据库快照",
                attempts=0,
                cached=True,
                duration_ms=0,
                source_table=price.get("source_table"),
                source_column=price.get("source_column"),
                source_time=price.get("source_time"),
                artifact_ref=artifact_ref,
            )
        return result, artifact_ref

    started = time.perf_counter()
    try:
        increment_metric(config, "redshift_batch_queries")
        if config.get("mock"):
            result = mock_database_results(manifest)
        else:
            from redshift_backend import enrich_products

            result = enrich_products(
                config["redshift_connection"], identities, config.get("redshift_schema") or None
            )
        duration_ms = round((time.perf_counter() - started) * 1000)
        artifact = store.write_artifact("redshift_snapshot", cache_key, result)
        cached_value = {"result": result, "artifact_ref": str(artifact)}
        store.cache_put("redshift", cache_key, cached_value)
        for product in manifest["products"]:
            key = identity_key(product["platform"], product["product_id"])
            record = result.get(key, {})
            price = record.get("price", {})
            issues = record.get("review_issues", [])
            store.event(
                run_id,
                "redshift",
                "review" if issues else "success",
                platform=product["platform"],
                product_id=product["product_id"],
                severity="warning" if issues else "info",
                message=(
                    "; ".join(str(issue.get("message", issue)) for issue in issues)
                    if issues
                    else f"platform_goods_key={record.get('key', {}).get('value')}"
                ),
                attempts=1,
                duration_ms=duration_ms,
                source_table=price.get("source_table"),
                source_column=price.get("source_column"),
                source_time=price.get("source_time"),
                artifact_ref=str(artifact),
            )
        progress(f"PostgreSQL 批量查询完成，耗时 {duration_ms / 1000:.1f} 秒")
        return result, str(artifact)
    except Exception as exc:
        message = safe_message(exc)
        progress(f"PostgreSQL 查询失败：{message}")
        for product in manifest["products"]:
            store.event(
                run_id,
                "redshift",
                "failed",
                platform=product["platform"],
                product_id=product["product_id"],
                severity="error",
                error_code="REDSHIFT_ERROR",
                message=message,
                attempts=1,
                duration_ms=round((time.perf_counter() - started) * 1000),
            )
        raise DemoError("REDSHIFT_ERROR", message) from exc


def attributes_map(record: dict[str, Any]) -> dict[str, str]:
    values: dict[str, str] = {}
    for attribute in record.get("attributes", []):
        name = str(attribute.get("name", "")).strip()
        value = str(attribute.get("value", "")).strip()
        if name and value and name not in values:
            values[name] = value
    return values


def normalized_attribute_name(value: Any) -> str:
    return re.sub(r"[\s:：()（）/_-]+", "", str(value or "")).casefold()


def database_enrichment(attrs: dict[str, str]) -> dict[str, dict[str, str] | None]:
    normalized = {
        normalized_attribute_name(name): (name, value)
        for name, value in attrs.items()
        if first_non_empty(value) is not None
    }
    resolved: dict[str, dict[str, str] | None] = {}
    for field, aliases in ENRICHMENT_ATTRIBUTE_ALIASES.items():
        resolved[field] = None
        for alias in aliases:
            match = normalized.get(normalized_attribute_name(alias))
            if match:
                resolved[field] = {"value": match[1], "attribute_name": match[0]}
                break
    return resolved


def valid_search_sources(sources: Any) -> list[dict[str, str]]:
    valid: list[dict[str, str]] = []
    for item in sources if isinstance(sources, list) else []:
        if not isinstance(item, dict):
            continue
        url = str(item.get("url") or "").strip()
        parsed = urllib.parse.urlparse(url)
        if parsed.scheme in {"http", "https"} and parsed.netloc:
            valid.append({"url": url, "title": str(item.get("title") or "")})
    return valid


def resolve_enrichment(
    model: dict[str, Any],
    attrs: dict[str, str],
    search: dict[str, Any],
    search_sources: Any,
) -> dict[str, Any]:
    database = database_enrichment(attrs)
    provider_sources = valid_search_sources(search_sources)
    search_has_evidence = bool(provider_sources)
    values: dict[str, Any] = {}
    selected_sources: dict[str, dict[str, Any]] = {}
    for field in (*ENRICHMENT_FIELDS, "hot_topics"):
        model_value = first_non_empty(model.get(field)) if field in ENRICHMENT_FIELDS else None
        database_match = database.get(field) if field in ENRICHMENT_FIELDS else None
        search_value = first_non_empty(search.get(field)) if search_has_evidence else None
        if model_value is not None:
            values[field] = model_value
            selected_sources[field] = {"source": "qwen_extract"}
        elif database_match:
            values[field] = database_match["value"]
            selected_sources[field] = {
                "source": "database.attributes",
                "attribute_name": database_match["attribute_name"],
            }
        elif search_value is not None:
            values[field] = search_value
            selected_sources[field] = {
                "source": "qwen_search",
                "provider_sources": provider_sources,
            }
        else:
            values[field] = None
            selected_sources[field] = {"source": "unresolved"}
    search_rejected = (
        any(first_non_empty(value) is not None for value in search.values())
        and not search_has_evidence
    )
    return {
        "values": values,
        "selected_sources": selected_sources,
        "missing_fields": [field for field, value in values.items() if value is None],
        "search_rejected": search_rejected,
        "provider_sources": provider_sources,
    }


def missing_search_fields(
    model: dict[str, Any],
    attrs: dict[str, str],
    *,
    include_hot_topics: bool,
) -> list[str]:
    database = database_enrichment(attrs)
    missing = [
        field
        for field in ENRICHMENT_FIELDS
        if first_non_empty(model.get(field)) is None and database.get(field) is None
    ]
    if include_hot_topics:
        missing.append("hot_topics")
    return missing


def attributes_prompt(record: dict[str, Any]) -> str:
    values = attributes_map(record)
    return "\n".join(f"- {name}: {value}" for name, value in sorted(values.items()))


def skipped_qwen(code: str, message: str) -> dict[str, Any]:
    return {
        "ok": False,
        "data": {},
        "duration_ms": 0,
        "attempts": 0,
        "raw_json_ref": None,
        "error": {"code": code, "message": message},
        "cached": False,
    }


def skipped_search(code: str, message: str, missing_fields: list[str]) -> dict[str, Any]:
    return {
        "ok": True,
        "data": {},
        "provider_sources": [],
        "duration_ms": 0,
        "attempts": 0,
        "raw_json_ref": None,
        "error": None,
        "cached": False,
        "skipped": True,
        "skip_code": code,
        "skip_message": message,
        "requested_fields": missing_fields,
    }


def run_qwen_stage(
    run_id: str,
    manifest: dict[str, Any],
    database: dict[str, dict[str, Any]],
    ocr: dict[str, list[dict[str, Any]]],
    template: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
    extract_meter: ConcurrencyMeter,
    search_meter: ConcurrencyMeter,
    total_meter: ConcurrencyMeter,
) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    extract_results: dict[str, dict[str, Any]] = {}
    search_results: dict[str, dict[str, Any]] = {}
    extract_jobs: dict[futures.Future[dict[str, Any]], str] = {}
    with futures.ThreadPoolExecutor(
        max_workers=QWEN_EXTRACT_WORKERS, thread_name_prefix="qwen-extract"
    ) as extract_executor:
        for product in manifest["products"]:
            key = identity_key(product["platform"], product["product_id"])
            image_results = ocr[key]
            record = database.get(key)
            successful_images = [item for item in image_results if item.get("ok")]
            if not record:
                reason = "数据库前置节点无该商品记录，未调用 Qwen"
                extract_results[key] = skipped_qwen("PREREQUISITE_FAILED", reason)
                search_results[key] = skipped_qwen("PREREQUISITE_FAILED", reason)
                store.event(
                    run_id,
                    "qwen",
                    "skipped",
                    platform=product["platform"],
                    product_id=product["product_id"],
                    severity="warning",
                    error_code="PREREQUISITE_FAILED",
                    message=reason,
                )
                continue
            name = str(record.get("name", {}).get("value") or product["product_id"])
            attrs = attributes_prompt(record)
            snapshot_hash = str(record.get("db_snapshot_hash") or stable_hash(record))
            if successful_images:
                if len(successful_images) != len(image_results):
                    store.event(
                        run_id,
                        "qwen_extract",
                        "partial",
                        platform=product["platform"],
                        product_id=product["product_id"],
                        severity="warning",
                        error_code="OCR_PARTIAL_INPUT",
                        message=f"仅使用 {len(successful_images)}/{len(image_results)} 张成功 OCR 生成待复核草稿",
                    )
                ocr_text = "\n\n".join(
                    f"## 图片 {item['image_name']}\n{item['markdown']}" for item in successful_images
                )
                extract_future = extract_executor.submit(
                    run_qwen_one,
                    run_id,
                    product,
                    "extract",
                    name,
                    ocr_text,
                    attrs,
                    snapshot_hash,
                    template,
                    config,
                    store,
                    extract_meter,
                    total_meter,
                )
                extract_jobs[extract_future] = key
            else:
                reason = "全部图片 OCR 失败，无法调用 Qwen 结构化提取"
                extract_results[key] = skipped_qwen("PREREQUISITE_FAILED", reason)
                store.event(
                    run_id,
                    "qwen_extract",
                    "skipped",
                    platform=product["platform"],
                    product_id=product["product_id"],
                    severity="warning",
                    error_code="PREREQUISITE_FAILED",
                    message=reason,
                )
        for future in futures.as_completed(extract_jobs):
            extract_results[extract_jobs[future]] = future.result()

    search_jobs: dict[futures.Future[dict[str, Any]], tuple[str, list[str]]] = {}
    include_hot_topics = bool(config.get("include_hot_topics", True))
    bulk_first_pass = bool(config.get("bulk_first_pass", False))
    with futures.ThreadPoolExecutor(
        max_workers=QWEN_SEARCH_WORKERS, thread_name_prefix="qwen-search"
    ) as search_executor:
        for product in manifest["products"]:
            key = identity_key(product["platform"], product["product_id"])
            if key in search_results:
                continue
            record = database.get(key)
            if not record:
                continue
            extract_result = extract_results.get(key, {})
            model = extract_result.get("data", {}) if extract_result.get("ok") else {}
            requested_fields = missing_search_fields(
                model,
                attributes_map(record),
                include_hot_topics=include_hot_topics,
            )
            if bulk_first_pass:
                code = "BULK_FIRST_PASS"
                message = "批量首轮已关闭联网搜索；仅输出 OCR、数据库和待补全字段"
            elif not template.get("search", {}).get("enabled", True):
                code = "SEARCH_DISABLED"
                message = "模板已关闭联网搜索"
            elif not requested_fields:
                code = "SEARCH_NOT_NEEDED"
                message = "OCR明确内容和数据库属性已覆盖目标字段，未调用联网搜索"
            else:
                name = str(record.get("name", {}).get("value") or product["product_id"])
                snapshot_hash = str(record.get("db_snapshot_hash") or stable_hash(record))
                search_future = search_executor.submit(
                    run_qwen_one,
                    run_id,
                    product,
                    "search",
                    name,
                    "",
                    "",
                    snapshot_hash,
                    template,
                    config,
                    store,
                    search_meter,
                    total_meter,
                )
                search_jobs[search_future] = (key, requested_fields)
                continue
            search_results[key] = skipped_search(code, message, requested_fields)
            store.event(
                run_id,
                "qwen_search",
                "skipped",
                platform=product["platform"],
                product_id=product["product_id"],
                severity="info",
                error_code=code,
                message=message,
            )
        for future in futures.as_completed(search_jobs):
            key, requested_fields = search_jobs[future]
            search_results[key] = {**future.result(), "requested_fields": requested_fields}
    return extract_results, search_results


def issue(code: str, message: str, severity: str = "error") -> dict[str, str]:
    return {"code": code, "message": message, "severity": severity}


def blue_hat_value(attrs: dict[str, str], model_value: Any) -> tuple[str | None, str | None]:
    mismatch = None
    for value in (attrs.get("批准文号"), attrs.get("蓝帽标识"), model_value):
        text = str(value or "").strip()
        if not text:
            continue
        compact = re.sub(r"\s+", "", text)
        if re.search(r"国食注字TY", compact, re.I):
            mismatch = text
            continue
        if re.search(r"(?:国食健字|卫食健字|食健备|国食健注|国食健备).*\d", compact):
            return text, mismatch
    return None, mismatch


def comparable_brand(value: Any) -> str:
    text = re.sub(r"[（(][^）)]*[）)]", "", str(value or "").casefold())
    return re.sub(r"[\W_]+", "", text)


def has_daily_dosage_conflict(value: Any) -> bool:
    doses = re.findall(r"(?:每日|每天)\s*[^、，；。\n]+", str(value or ""))
    return len(doses) > 1


def has_competitor_image_note(value: Any) -> bool:
    return bool(re.search(r"竞品|非本商品|不属于本商品|品牌不一致|勿混淆", str(value or "")))


def resolve_extra_output(column: dict[str, Any], context: dict[str, Any]) -> Any:
    """Resolve simple new template fields without turning the Demo into a workflow engine."""
    values: list[Any] = []
    for source in column.get("sources", []):
        prefix, _, key = str(source).partition(".")
        source_object = context.get(prefix)
        if isinstance(source_object, dict) and key:
            current: Any = source_object
            for part in key.split("."):
                current = current.get(part) if isinstance(current, dict) else None
            values.append(current)
    if column.get("strategy") == "json_serialize":
        return json.dumps(context.get("model", {}), ensure_ascii=False) or None
    return first_non_empty(*values)


def assemble_product(
    product: dict[str, Any],
    db_record: dict[str, Any] | None,
    db_artifact_ref: str,
    image_results: list[dict[str, Any]],
    extract_result: dict[str, Any],
    search_result: dict[str, Any],
    template: dict[str, Any],
) -> dict[str, Any]:
    problems: list[dict[str, str]] = []
    if db_record is None:
        db_record = {}
        problems.append(issue("DATABASE_RECORD_MISSING", "数据库批量结果缺少该商品"))
    for item in db_record.get("review_issues", []):
        problems.append(
            issue(
                str(item.get("code", "DATABASE_REVIEW")),
                str(item.get("message", item)),
                str(item.get("severity", "error")),
            )
        )
    for image in image_results:
        if not image.get("ok"):
            error = image.get("error", {})
            problems.append(
                issue(
                    str(error.get("code", "OCR_FAILED")),
                    f"图片 {image['image_name']} OCR 最终失败: {error.get('message', '')}",
                )
            )
    if not extract_result.get("ok"):
        error = extract_result.get("error", {})
        problems.append(issue(str(error.get("code", "QWEN_EXTRACT_FAILED")), str(error.get("message", "Qwen 提取失败"))))
    if not search_result.get("ok") and search_result.get("error", {}).get("code") != "PREREQUISITE_FAILED":
        error = search_result.get("error", {})
        problems.append(issue(str(error.get("code", "SEARCH_FAILED")), str(error.get("message", "联网搜索失败")), "warning"))

    identity = db_record.get("identity", {})
    db_name = db_record.get("name", {})
    db_price = db_record.get("price", {})
    if identity.get("platform") not in (None, product["platform"]) or identity.get("product_id") not in (None, product["product_id"]):
        problems.append(issue("PROTECTED_IDENTITY_MISMATCH", "数据库返回身份与所选 (platform, product_id) 不一致"))
    if not db_name.get("value"):
        problems.append(issue("NAME_NOT_FOUND", "没有可用商品名称"))
    if db_price.get("status") != "ok" or positive_decimal(db_price.get("value")) is None:
        if not any(item["code"].startswith("PRICE_") for item in problems):
            problems.append(issue("PRICE_NOT_FOUND", "没有可用数据库价格"))
    price_month, price_freshness = price_month_and_freshness(db_price.get("source_time"))
    if price_month and str(db_price.get("source_table") or "").startswith("mv_com_goods_statistics_monthly"):
        year, month = map(int, price_month.split("-"))
        age_months = date.today().year * 12 + date.today().month - (year * 12 + month)
        if age_months > 3:
            problems.append(
                issue(
                    "PRICE_STALE",
                    f"月度价格来源 {price_month} 距当前已 {age_months} 个月，请人工确认",
                    "warning",
                )
            )

    model = extract_result.get("data", {}) if extract_result.get("ok") else {}
    if has_daily_dosage_conflict(model.get("ri_fu_liang")):
        problems.append(
            issue(
                "DOSAGE_CONFLICT",
                f"日服量存在多个相互独立的每日用量：{model.get('ri_fu_liang')}",
            )
        )
    if has_competitor_image_note(model.get("notes")):
        problems.append(
            issue(
                "IMAGE_PRODUCT_MISMATCH",
                "Qwen 提示存在疑似竞品或非本商品图片，请人工确认后再使用营销字段。",
            )
        )
    attrs = attributes_map(db_record)
    search_data = search_result.get("data", {}) if search_result.get("ok") else {}
    enrichment = resolve_enrichment(
        model,
        attrs,
        search_data,
        search_result.get("provider_sources", []),
    )
    search_sources = enrichment["provider_sources"]
    if enrichment["search_rejected"]:
        problems.append(issue("SEARCH_EVIDENCE_MISSING", "联网搜索没有返回可核验的提供方来源，未采用搜索值", "warning"))
    search = search_data if search_sources else {}
    parsed = parse_spec(model.get("guige"))
    parsed_total = positive_decimal(parsed.get("total"))
    model_total = positive_decimal(model.get("guige_zong_liang"))
    if parsed_total is not None and model_total is not None and parsed_total != model_total:
        problems.append(
            issue(
                "SPEC_TOTAL_CONFLICT",
                f"确定性规格总量 {decimal_text(parsed_total)} 与模型总量 {decimal_text(model_total)} 不一致",
            )
        )
    total = parsed.get("total") if parsed_total is not None else model.get("guige_zong_liang")
    if parsed.get("status") == "unparsed" and model_total is not None:
        problems.append(issue("SPEC_PARSE_FALLBACK", "规格无法确定性解析，暂用模型总量", "warning"))
    costs = calculate_costs(
        db_price.get("value"),
        total,
        model.get("min_ri_fu_liang"),
        model.get("max_ri_fu_liang"),
        parsed.get("unit"),
        model.get("ri_fu_liang"),
    )
    if costs["status"] == "unit_mismatch":
        problems.append(issue("DOSAGE_UNIT_MISMATCH", "规格单位与日服量单位不一致，未计算日服成本", "warning"))
    database_brand = attrs.get("品牌")
    model_brand = str(model.get("brand") or "").strip() or None
    database_brand_key = comparable_brand(database_brand)
    model_brand_key = comparable_brand(model_brand)
    if (
        database_brand
        and model_brand
        and database_brand_key
        and model_brand_key
        and database_brand_key not in model_brand_key
        and model_brand_key not in database_brand_key
    ):
        problems.append(
            issue(
                "BRAND_SOURCE_CONFLICT",
                f"数据库品牌“{database_brand}”与OCR提取品牌“{model_brand}”不一致，结果按数据库优先",
                "warning",
            )
        )
    blue_hat, regulatory_mismatch = blue_hat_value(attrs, model.get("blue_hat"))
    if regulatory_mismatch:
        problems.append(
            issue(
                "BLUE_HAT_TYPE_MISMATCH",
                f"“{regulatory_mismatch}”属于特医食品注册号，不作为蓝帽标识输出",
                "warning",
            )
        )
    persona_error = ""
    if not extract_result.get("ok"):
        persona_error = str(extract_result.get("error", {}).get("message", "Qwen 提取失败"))
    all_skus = product.get("all_skus")
    sku_values = list(dict.fromkeys(str(value).strip() for value in all_skus if str(value).strip())) if isinstance(all_skus, list) else []
    available: dict[str, Any] = {
        "platform": product["platform"],
        "product_id": product["product_id"],
        "商品名称": db_name.get("value"),
        "价格": as_decimal(db_price.get("value")),
        "规格": parsed.get("main_spec"),
        "包装": first_non_empty(
            parsed.get("package"),
            attrs.get("包装规格"),
            attrs.get("包装形式"),
            attrs.get("包装"),
            attrs.get("包装清单"),
        ),
        "规格总量": total,
        "最小单位价格": costs.get("unit_price"),
        "是否多规格": model.get("is_multi_spec"),
        "日服量": model.get("ri_fu_liang"),
        "最小日服量": model.get("min_ri_fu_liang"),
        "最大日服量": model.get("max_ri_fu_liang"),
        "最小日服成本": costs.get("min_daily_cost"),
        "最大日服成本": costs.get("max_daily_cost"),
        "成分": model.get("ingredients"),
        "人群": first_non_empty(model.get("renqun"), model.get("marketing_crowd")),
        "功能": model.get("gongneng"),
        "推送标题": model.get("push_title"),
        "主打卖点": model.get("main_selling_points"),
        "推送时间": model.get("push_time"),
        "品牌": first_non_empty(database_brand, model_brand),
        "剂型": first_non_empty(attrs.get("剂型"), model.get("dosage_form")),
        "蓝帽标识": blue_hat,
        "代工厂": enrichment["values"]["factory_name"],
        "代工厂地址": enrichment["values"]["factory_address"],
        "生产许可证": enrichment["values"]["production_license"],
        "产地": enrichment["values"]["origin"],
        "热门话题": enrichment["values"]["hot_topics"],
        "成分含量": model.get("ingredient_content"),
        "适用人群": first_non_empty(model.get("applicable_crowd"), model.get("renqun")),
        "商品主数据来源": db_record.get("platform_source", {}).get("label"),
        "价格数据来源": db_price.get("platform_source", {}).get("label"),
        "价格月份": price_month,
        "价格新鲜度": price_freshness,
        "所有SKU": "、".join(sku_values) if sku_values else None,
        "SKU数量": len(sku_values) if sku_values else None,
        "上架时间": product.get("first_shelf_time"),
        "persona_json": json.dumps(model, ensure_ascii=False) if model else None,
        "persona_error": persona_error or None,
    }
    context = {
        "identity": {"platform": product["platform"], "product_id": product["product_id"]},
        "database": {**db_record, "product_name": db_name},
        "attributes": attrs,
        "model": model,
        "search": search,
        "input": product,
    }
    row = {
        "platform": product["platform"],
        "商品主数据来源": available["商品主数据来源"],
        "价格数据来源": available["价格数据来源"],
        "价格月份": available["价格月份"],
        "价格新鲜度": available["价格新鲜度"],
    }
    for column in template["output_columns"]:
        name = column["name"]
        row[name] = available[name] if name in available else resolve_extra_output(column, context)
    blocking = [item for item in problems if item["severity"] == "error"]
    attribute_evidence = {
        str(item.get("name")): item for item in db_record.get("attributes", [])
    }
    source_evidence: dict[str, Any] = {
        "platform": {"selected": "identity.platform", "value": product["platform"]}
    }
    for column in template["output_columns"]:
        name = column["name"]
        if name == "product_id":
            selected: Any = {"source": "identity.product_id", "value": product["product_id"]}
        elif name == "商品名称":
            selected = {"source": db_name.get("source_table"), "time": db_name.get("source_time")}
        elif name == "价格":
            selected = {
                "source": db_price.get("source_table"),
                "column": db_price.get("source_column"),
                "time": db_price.get("source_time"),
                "raw_value": db_price.get("raw_value"),
            }
        elif name in {"品牌", "剂型"} and attrs.get(name):
            selected = attribute_evidence.get(name, {"source": "database.attributes"})
        elif name == "蓝帽标识" and blue_hat:
            selected = attribute_evidence.get("批准文号") or attribute_evidence.get("蓝帽标识") or {
                "source": "qwen_extract"
            }
        elif name in ENRICHMENT_OUTPUT_NAMES:
            field = ENRICHMENT_OUTPUT_NAMES[name]
            selected = dict(enrichment["selected_sources"][field])
            attribute_name = selected.get("attribute_name")
            if attribute_name and attribute_name in attribute_evidence:
                selected["attribute_evidence"] = attribute_evidence[attribute_name]
            if selected.get("source") == "qwen_extract":
                selected["artifact"] = extract_result.get("raw_json_ref")
            elif selected.get("source") == "qwen_search":
                selected["artifact"] = search_result.get("raw_json_ref")
        elif column.get("strategy") == "derived" or any(
            str(source).startswith("derived.") for source in column.get("sources", [])
        ):
            selected = {"source": "deterministic_python", "inputs": column.get("sources", [])}
        elif any(str(source).startswith("search.") for source in column.get("sources", [])):
            selected = {
                "source": "qwen_search",
                "artifact": search_result.get("raw_json_ref"),
                "provider_sources": search_sources,
            }
        elif any(str(source).startswith("model.") for source in column.get("sources", [])):
            selected = {"source": "qwen_extract", "artifact": extract_result.get("raw_json_ref")}
        else:
            selected = {"source": "local_input_or_first_non_empty"}
        source_evidence[column["name"]] = {
            "declared_sources": column.get("sources", []),
            "protected": bool(column.get("protected")),
            "selected": selected,
        }
    return {
        "identity": {"platform": product["platform"], "product_id": product["product_id"]},
        "status": "review" if blocking else "success",
        "fields": row,
        "database": db_record,
        "source_evidence": source_evidence,
        "missing_enrichment_fields": enrichment["missing_fields"],
        "artifacts": {
            "database": db_artifact_ref,
            "ocr": [
                {
                    "image": item["image_name"],
                    "sha256": item["sha256"],
                    "log_id": item.get("log_id"),
                    "raw_json": item.get("raw_json_ref"),
                    "markdown": item.get("markdown_ref"),
                }
                for item in image_results
            ],
            "qwen_extract": extract_result.get("raw_json_ref"),
            "qwen_search": search_result.get("raw_json_ref"),
            "qwen_search_sources": search_sources,
        },
        "validation_issues": problems,
    }


AUDIT_COLUMNS = [
    ("run_id", "运行 ID", "text", 24),
    ("attempt_id", "尝试 ID", "text", 28),
    ("platform", "平台", "text", 12),
    ("product_id", "商品 ID", "text", 18),
    ("stage", "节点", "text", 18),
    ("status", "状态", "text", 12),
    ("severity", "级别", "text", 12),
    ("error_code", "错误码", "text", 24),
    ("message", "说明", "text", 52),
    ("attempts", "尝试次数", "integer", 12),
    ("cached", "命中缓存", "boolean", 12),
    ("duration_ms", "耗时(ms)", "integer", 14),
    ("source_table", "来源表", "text", 36),
    ("source_column", "来源字段", "text", 26),
    ("source_time", "来源时间", "datetime", 22),
    ("artifact_ref", "产物引用", "text", 52),
    ("created_at", "记录时间", "datetime", 22),
]


def workbook_columns(template: dict[str, Any]) -> list[dict[str, Any]]:
    type_map = {
        "string": "text",
        "decimal": "number",
        "number": "number",
        "integer": "integer",
        "boolean": "boolean",
        "datetime": "datetime",
        "array": "json",
    }
    columns = [
        {"key": "platform", "header": "platform", "type": "text", "width": 12},
        {"key": "商品主数据来源", "header": "商品主数据来源", "type": "text", "width": 28},
        {"key": "价格数据来源", "header": "价格数据来源", "type": "text", "width": 28},
        {"key": "价格月份", "header": "价格月份", "type": "text", "width": 12},
        {"key": "价格新鲜度", "header": "价格新鲜度", "type": "text", "width": 20},
    ]
    for item in template["output_columns"]:
        field_type = type_map[item["type"]]
        column: dict[str, Any] = {
            "key": item["name"],
            "header": item["name"],
            "type": field_type,
        }
        if field_type == "number":
            column["format"] = "#,##0.0000" if "成本" in item["name"] or "单位价格" in item["name"] else "#,##0.00"
        if item["name"] in {"商品名称", "persona_json", "persona_error", "主打卖点"}:
            column["width"] = 42
        elif item["name"] == "product_id":
            column["width"] = 18
        columns.append(column)
    return columns


def workbook_payload(
    run_id: str,
    template: dict[str, Any],
    documents: list[dict[str, Any]],
    events: list[dict[str, Any]],
) -> dict[str, Any]:
    audit_rows = []
    for event in events:
        row = dict(event)
        source_time = row.get("source_time")
        compact = str(source_time).strip() if source_time is not None else ""
        if re.fullmatch(r"\d{6}", compact):
            row["source_time"] = f"{compact[:4]}-{compact[4:]}-01T00:00:00Z"
        elif re.fullmatch(r"\d{8}", compact):
            row["source_time"] = f"{compact[:4]}-{compact[4:6]}-{compact[6:]}T00:00:00Z"
        audit_rows.append(row)
    return {
        "schema_version": template["version"],
        "run": {
            "run_id": run_id,
            "generated_at": iso_now(),
            "template_id": template["template_id"],
            "template_version": template["version"],
        },
        "result": {
            "columns": workbook_columns(template),
            "rows": [document["fields"] for document in documents],
        },
        "audit": {
            "columns": [
                {"key": key, "header": header, "type": item_type, "width": width}
                for key, header, item_type, width in AUDIT_COLUMNS
            ],
            "rows": audit_rows,
        },
    }


def export_csv_file(output_dir: Path, template: dict[str, Any], documents: list[dict[str, Any]]) -> Path:
    columns = workbook_columns(template)
    path = output_dir / "result.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[column["key"] for column in columns], extrasaction="ignore")
        writer.writeheader()
        for document in documents:
            row: dict[str, Any] = {}
            for column in columns:
                value = document["fields"].get(column["key"])
                row[column["key"]] = json.dumps(value, ensure_ascii=False, default=json_default) if isinstance(value, (list, dict)) else value
            writer.writerow(row)
    return path


def export_workbook_file(
    output_dir: Path,
    payload: dict[str, Any],
    _config: dict[str, Any],
) -> Path:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise DemoError("OPENPYXL_NOT_FOUND", "缺少 openpyxl，请通过启动器自动安装") from exc

    def cell_value(value: Any, item_type: str) -> Any:
        if value is None:
            return None
        if item_type == "number":
            number = as_decimal(value)
            return float(number) if number is not None else None
        if item_type == "integer":
            number = as_decimal(value)
            return int(number) if number is not None else None
        if item_type == "boolean":
            return bool(value)
        if item_type == "datetime":
            try:
                parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
                return parsed.replace(tzinfo=None)
            except ValueError:
                return str(value)
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=json_default)
        return "'" + text if text.startswith(("=", "+", "-", "@")) else text

    def add_sheet(workbook: Any, name: str, section: dict[str, Any]) -> None:
        sheet = workbook.create_sheet(name)
        columns = section["columns"]
        header_fill = PatternFill("solid", fgColor="1F4E78")
        stripe_fill = PatternFill("solid", fgColor="DDEBF7")
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(1, index, column["header"])
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            column_letter = get_column_letter(index)
            sheet.column_dimensions[column_letter].width = min(float(column.get("width", 16)), 52)
            if name == "Result" and column.get("key") == "persona_json":
                sheet.column_dimensions[column_letter].hidden = True
        for row_index, row in enumerate(section["rows"], 2):
            for column_index, column in enumerate(columns, 1):
                cell = sheet.cell(
                    row_index,
                    column_index,
                    cell_value(row.get(column["key"]), column.get("type", "text")),
                )
                if row_index % 2 == 0:
                    cell.fill = stripe_fill
                wrap = column.get("key") in {"商品名称", "主打卖点", "message"}
                cell.alignment = Alignment(vertical="top", wrap_text=wrap)
                if column.get("format"):
                    cell.number_format = column["format"]
                elif column.get("type") == "integer":
                    cell.number_format = "#,##0"
                elif column.get("type") == "datetime":
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.freeze_panes = "C2" if name == "Result" else "D2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        for row_index in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_index].height = 30 if name == "Result" else 24

    target = output_dir / "result.xlsx"
    temporary = target.with_suffix(".xlsx.tmp")
    try:
        workbook = Workbook()
        workbook.remove(workbook.active)
        add_sheet(workbook, "Result", payload["result"])
        add_sheet(workbook, "Audit", payload["audit"])
        workbook.save(temporary)
        temporary.replace(target)
        verified = load_workbook(target, read_only=True, data_only=False)
        try:
            if verified.sheetnames != ["Result", "Audit"]:
                raise DemoError("WORKBOOK_EXPORT_FAILED", "Excel 工作表结构校验失败")
            if verified["Result"].max_column != len(payload["result"]["columns"]):
                raise DemoError("WORKBOOK_EXPORT_FAILED", "Result 列数校验失败")
        finally:
            verified.close()
        return target
    except DemoError:
        raise
    except Exception as exc:
        raise DemoError("WORKBOOK_EXPORT_FAILED", safe_message(exc)) from exc
    finally:
        temporary.unlink(missing_ok=True)


def write_partial_product(
    run_id: str,
    output_dir: Path,
    products_dir: Path,
    product: dict[str, Any],
    completed_images: list[dict[str, Any]],
    pending_images: list[str],
    database: dict[str, dict[str, Any]],
    db_artifact_ref: str | None,
    template: dict[str, Any],
    config: dict[str, Any],
    store: StateStore,
    extract_meter: ConcurrencyMeter,
    search_meter: ConcurrencyMeter,
    total_meter: ConcurrencyMeter,
    partial_documents: dict[str, dict[str, Any]],
    partial_lock: threading.Lock,
) -> None:
    key = identity_key(product["platform"], product["product_id"])
    # 部分结果只做结构化提取，联网搜索留到最终完整结果，避免重复外发和重复等待。
    partial_config = dict(config)
    partial_config["bulk_first_pass"] = True
    partial_config["include_hot_topics"] = False
    extract_results, search_results = run_qwen_stage(
        run_id,
        {"products": [product]},
        database,
        {key: completed_images},
        template,
        partial_config,
        store,
        extract_meter,
        search_meter,
        total_meter,
    )
    document = assemble_product(
        product,
        database.get(key),
        db_artifact_ref,
        completed_images,
        extract_results.get(key, skipped_qwen("MISSING", "缺少部分提取结果")),
        search_results.get(key, skipped_search("BULK_FIRST_PASS", "部分结果暂不执行联网搜索", [])),
        template,
    )
    if pending_images:
        document["validation_issues"].append(
            {
                "code": "OCR_TIMEOUT_REVIEW",
                "message": (
                    f"已使用 {len(completed_images)} 张已完成 OCR 图片生成临时结果；"
                    f"图片 {', '.join(pending_images)} 超过 30 秒上限，已放弃重试并标记待复核"
                ),
                "severity": "warning",
            }
        )
        document["status"] = "review"
    with partial_lock:
        partial_documents[key] = document
        documents = list(partial_documents.values())
    dump_json(products_dir / f"{product['product_id']}.json", document)
    if pending_images:
        store.event(
            run_id,
            "validation",
            "review",
            platform=product["platform"],
            product_id=product["product_id"],
            severity="warning",
            error_code="OCR_TIMEOUT_REVIEW",
            message=document["validation_issues"][-1]["message"],
            artifact_ref=str(products_dir / f"{product['product_id']}.json"),
        )
    write_report(
        output_dir,
        run_id,
        "review" if any(item["status"] != "success" for item in documents) else "running",
        documents,
        store.events(run_id),
        config.get("metrics", {}),
        {
            "ocr": 0,
            "ocr_limit": ocr_worker_count(config),
            "qwen_extract": extract_meter.peak,
            "qwen_search": search_meter.peak,
            "qwen_total": total_meter.peak,
        },
        filename="实时结果.html",
        auto_refresh_seconds=3,
    )
    if pending_images:
        progress(
            f"实时结果已更新：{product['product_id']} 已用 {len(completed_images)} 张图片进入 Qwen，"
            f"图片 {', '.join(pending_images)} 已标记待复核；文件：{output_dir / '实时结果.html'}"
        )
    else:
        progress(
            f"实时结果已更新：{product['product_id']} 已完成 OCR 并已进入 Qwen；"
            f"文件：{output_dir / '实时结果.html'}"
        )


def write_report(
    output_dir: Path,
    run_id: str,
    run_status: str,
    documents: list[dict[str, Any]],
    events: list[dict[str, Any]],
    metrics: dict[str, Any],
    peaks: dict[str, int],
    performance: dict[str, Any] | None = None,
    filename: str = "report.html",
    auto_refresh_seconds: int = 0,
) -> Path:
    success_count = sum(document["status"] == "success" for document in documents)
    review_count = len(documents) - success_count
    cache_hits = sum(bool(event.get("cached")) for event in events)
    error_events = [
        event for event in events
        if event.get("severity") == "error" or event.get("status") == "failed"
    ]
    error_rows = "".join(
        "<tr>"
        f"<td>{html.escape(str(event.get('product_id') or '全局'))}</td>"
        f"<td>{html.escape(str(event.get('stage') or ''))}</td>"
        f"<td>{html.escape(str(event.get('error_code') or ''))}</td>"
        f"<td>{html.escape(str(event.get('message') or ''))}</td>"
        "</tr>"
        for event in error_events
    )
    error_section = (
        "<section><h2>失败原因</h2>"
        "<table><thead><tr><th>商品</th><th>节点</th><th>错误码</th><th>说明</th></tr></thead>"
        f"<tbody>{error_rows}</tbody></table></section>"
        if error_rows else ""
    )
    performance_section = ""
    if performance:
        ocr_performance = performance.get("ocr", {})
        stages = performance.get("stages_ms", {})
        comparison = performance.get("comparison_to_previous")
        comparison_rows = ""
        if isinstance(comparison, dict):
            comparison_rows = (
                f"<tr><td>相较上次总耗时</td><td>{comparison.get('total_ms_delta', '-')} ms "
                f"({comparison.get('total_ms_change_percent', '-')}%)</td></tr>"
                f"<tr><td>相较上次 OCR 吞吐</td><td>{comparison.get('ocr_success_images_per_minute_delta', '-')} 张/分钟</td></tr>"
            )
        performance_section = f"""
        <section><h2>性能摘要</h2>
        <table><thead><tr><th>指标</th><th>值</th></tr></thead><tbody>
        <tr><td>总耗时</td><td>{performance.get('total_ms', 0)} ms</td></tr>
        <tr><td>PostgreSQL 耗时</td><td>{stages.get('redshift_ms', 0)} ms</td></tr>
        <tr><td>OCR 耗时</td><td>{stages.get('ocr_ms', 0)} ms</td></tr>
        <tr><td>PostgreSQL + OCR 并行墙钟耗时</td><td>{stages.get('parallel_wall_ms', 0)} ms</td></tr>
        <tr><td>Qwen 阶段耗时</td><td>{stages.get('qwen_ms', 0)} ms</td></tr>
        <tr><td>OCR P50 / P95（含超时失败）</td><td>{ocr_performance.get('p50_ms', '-')} / {ocr_performance.get('p95_ms', '-')} ms</td></tr>
        <tr><td>OCR P50 / P95（仅成功图片）</td><td>{ocr_performance.get('success_p50_ms', '-')} / {ocr_performance.get('success_p95_ms', '-')} ms</td></tr>
        <tr><td>OCR 吞吐</td><td>{ocr_performance.get('success_images_per_minute', '-')} 张/分钟</td></tr>
        <tr><td>OCR 图片：计划 / 已完成 / 未开始</td><td>{ocr_performance.get('planned_images', 0)} / {ocr_performance.get('completed_images', 0)} / {ocr_performance.get('unstarted_images', 0)}</td></tr>
        <tr><td>OCR 输入压缩</td><td>{ocr_performance.get('input_source_bytes', 0)} → {ocr_performance.get('input_sent_bytes', 0)} bytes；缩放 {ocr_performance.get('resized_images', 0)} 张</td></tr>
        <tr><td>OCR 成功 / 失败 / 缓存</td><td>{ocr_performance.get('success_images', 0)} / {ocr_performance.get('failed_images', 0)} / {ocr_performance.get('cache_hits', 0)}</td></tr>
        {comparison_rows}
        </tbody></table><p class="sub">完整机器可读数据：performance.json</p></section>
        """
    product_sections = []
    for document in documents:
        identity = document["identity"]
        product_id = identity["product_id"]
        product_events = [
            event
            for event in events
            if event.get("platform") == identity["platform"]
            and event.get("product_id") == product_id
        ]
        event_rows = "".join(
            "<tr>"
            f"<td>{html.escape(str(event.get('stage') or ''))}</td>"
            f"<td><span class='status {html.escape(str(event.get('status') or ''))}'>{html.escape(str(event.get('status') or ''))}</span></td>"
            f"<td>{html.escape(str(event.get('duration_ms') if event.get('duration_ms') is not None else ''))}</td>"
            f"<td>{'是' if event.get('cached') else '否'}</td>"
            f"<td>{html.escape(str(event.get('message') or ''))}</td>"
            "</tr>"
            for event in product_events
        )
        issues = document.get("validation_issues", [])
        issue_html = "<ul>" + "".join(
            f"<li class='{html.escape(item['severity'])}'>{html.escape(item['code'])}: {html.escape(item['message'])}</li>"
            for item in issues
        ) + "</ul>" if issues else "<p>无校验问题</p>"
        price = document.get("database", {}).get("price", {})
        search_sources = document.get("artifacts", {}).get("qwen_search_sources", [])
        search_evidence = (
            "<p>搜索证据：" + "；".join(
                f"<a href='{html.escape(str(item.get('url') or ''))}'>{html.escape(str(item.get('title') or item.get('url') or '来源'))}</a>"
                for item in search_sources
            ) + "</p>"
            if search_sources else "<p>搜索证据：无</p>"
        )
        json_link = urllib.parse.quote(f"products/{product_id}.json")
        product_sections.append(
            f"""
            <section>
              <h2>{html.escape(identity['platform'])} / {html.escape(product_id)}
                <span class='status {document['status']}'>{document['status']}</span>
              </h2>
              <p>名称：{html.escape(str(document['fields'].get('商品名称') or '-'))}<br>
                 价格：{html.escape(str(document['fields'].get('价格') or '-'))}；
                 来源：{html.escape(str(price.get('source_table') or '-'))}.{html.escape(str(price.get('source_column') or '-'))}；
                 时间：{html.escape(str(price.get('source_time') or '-'))}<br>
                  <a href='{json_link}'>查看商品 JSON</a></p>
              {search_evidence}
              {issue_html}
              <table><thead><tr><th>节点</th><th>状态</th><th>耗时(ms)</th><th>缓存</th><th>说明</th></tr></thead>
              <tbody>{event_rows}</tbody></table>
            </section>
            """
        )
    generated = html.escape(iso_now())
    attempt_id = next((str(event.get("attempt_id")) for event in events if event.get("attempt_id")), "-")
    refresh_tag = (
        f'<meta http-equiv="refresh" content="{auto_refresh_seconds}">'
        if auto_refresh_seconds > 0
        else ""
    )
    body = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">{refresh_tag}<title>本地 Qwen-VL OCR Demo v5 运行报告</title>
<style>
body{{font-family:"Segoe UI","Microsoft YaHei",sans-serif;background:#f4f7fb;color:#203040;margin:0;padding:24px}}
main{{max-width:1280px;margin:auto}}h1{{margin:0 0 8px}}.sub{{color:#60758a}}
.cards{{display:flex;gap:12px;flex-wrap:wrap;margin:20px 0}}.card{{background:white;border-radius:10px;padding:14px 18px;min-width:145px;box-shadow:0 2px 8px #dbe4ee}}
section{{background:white;border-radius:10px;padding:18px;margin:16px 0;box-shadow:0 2px 8px #dbe4ee}}
table{{border-collapse:collapse;width:100%;font-size:13px}}th{{background:#1f4e78;color:white;text-align:left}}th,td{{padding:8px;border-bottom:1px solid #d9e2f3;vertical-align:top}}
.status{{display:inline-block;padding:2px 8px;border-radius:999px;background:#e7edf4;font-size:12px}}.success,.complete{{background:#dff3e7;color:#176b3a}}.review,.warning,.skipped{{background:#fff0c9;color:#895b00}}.failed,.error{{background:#fde2e2;color:#9d2020}}li.warning{{color:#895b00}}li.error{{color:#9d2020}}a{{color:#1663a5}}
</style></head><body><main>
<h1>本地 Qwen-VL OCR Demo v5</h1><div class="sub">run_id={html.escape(run_id)} · attempt_id={html.escape(attempt_id)} · {generated}</div>
<div class="cards">
  <div class="card"><b>运行状态</b><br>{html.escape(run_status)}</div>
  <div class="card"><b>成功商品</b><br>{success_count}</div>
  <div class="card"><b>待复核</b><br>{review_count}</div>
  <div class="card"><b>缓存命中</b><br>{cache_hits}</div>
  <div class="card"><b>OCR 峰值并发</b><br>{peaks.get('ocr', 0)} / {peaks.get('ocr_limit', OCR_WORKERS)}</div>
  <div class="card"><b>Qwen 峰值并发</b><br>{peaks.get('qwen_total', 0)} / {QWEN_EXTRACT_WORKERS + QWEN_SEARCH_WORKERS}</div>
</div>
<p class="sub">实际调用计数：{html.escape(json.dumps(metrics, ensure_ascii=False, sort_keys=True))}</p>
{performance_section}
{error_section}
{''.join(product_sections)}
</main></body></html>"""
    report = output_dir / filename
    report.write_text(body, encoding="utf-8")
    return report


def execute_pipeline(
    manifest: dict[str, Any],
    template_path: Path,
    state_dir: Path,
    output_root: Path,
    config: dict[str, Any],
    *,
    force_new: bool = False,
    interrupt_after_parallel: bool = False,
    skip_workbook: bool = False,
) -> dict[str, Any]:
    pipeline_started = time.perf_counter()
    template = load_json(template_path)
    store = StateStore(state_dir)
    documents: list[dict[str, Any]] = []
    partial_documents: dict[str, dict[str, Any]] = {}
    partial_lock = threading.Lock()
    partial_export_lock = threading.Lock()
    run_id = ""
    attempt_id = ""
    run_fingerprint = ""
    output_dir = output_root
    ocr_meter = ConcurrencyMeter()
    extract_meter = ConcurrencyMeter()
    search_meter = ConcurrencyMeter()
    qwen_total_meter = ConcurrencyMeter()
    try:
        store.cleanup_expired()
        run_fingerprint = stable_hash(
            {
                "manifest_hash": manifest["manifest_hash"],
                "template_version": template["version"],
                "bulk_first_pass": bool(config.get("bulk_first_pass", False)),
                "include_hot_topics": bool(config.get("include_hot_topics", True)),
            }
        )
        run_id, output_dir, resumed = store.choose_run(
            run_fingerprint, output_root, force_new=force_new
        )
        attempt_id = store.begin_attempt()
        products_dir = output_dir / "products"
        products_dir.mkdir(parents=True, exist_ok=True)
        config["run_output_dir"] = str(output_dir)
        if resumed:
            print(f"恢复未完成运行: {run_id}")
        if not config.get("mock") and config.get("ocr_provider", "vllm-qwen-vl") == "vllm-qwen-vl":
            verify_vllm_available(
                config["vllm_ocr_api_base"], config["vllm_ocr_api_key"]
            )
        progress(f"运行开始：run_id={run_id}，attempt_id={attempt_id}")
        progress("PostgreSQL 与 OCR 并行启动；数据库失败时仍保留已完成 OCR 产物")
        def timed(call: Callable[[], Any]) -> tuple[Any, int]:
            started = time.perf_counter()
            value = call()
            return value, round((time.perf_counter() - started) * 1000)

        database_future: futures.Future[tuple[Any, int]]

        def on_product_partial(
            product: dict[str, Any],
            completed_images: list[dict[str, Any]],
            pending_images: list[str],
        ) -> None:
            with partial_export_lock:
                (database_value, _database_ms) = database_future.result()
                partial_database, partial_db_artifact = database_value
                write_partial_product(
                    run_id,
                    output_dir,
                    products_dir,
                    product,
                    completed_images,
                    pending_images,
                    partial_database,
                    partial_db_artifact,
                    template,
                    config,
                    store,
                    extract_meter,
                    search_meter,
                    qwen_total_meter,
                    partial_documents,
                    partial_lock,
                )

        parallel_started = time.perf_counter()
        with futures.ThreadPoolExecutor(max_workers=2, thread_name_prefix="preflight") as preflight:
            database_future = preflight.submit(
                timed, lambda: run_database_stage(run_id, manifest, config, store)
            )
            ocr_future = preflight.submit(
                timed,
                lambda: (run_paddle_ocr_stage if config.get("ocr_provider") == "paddleocr-vl" else run_ocr_stage)(
                    run_id,
                    manifest,
                    config,
                    store,
                    ocr_meter,
                    on_product_partial=on_product_partial,
                ),
            )
            (database_value, database_ms), (ocr_results, ocr_ms) = (
                database_future.result(),
                ocr_future.result(),
            )
        database, db_artifact_ref = database_value
        stage_timings = {
            "redshift_ms": database_ms,
            "ocr_ms": ocr_ms,
            "parallel_wall_ms": round((time.perf_counter() - parallel_started) * 1000),
        }
        progress("PostgreSQL 与 OCR 前置阶段完成，开始 Qwen")
        if interrupt_after_parallel:
            raise DemoError("SIMULATED_INTERRUPT", "已在 OCR/Redshift 后模拟中断")
        qwen_started = time.perf_counter()
        extract_results, search_results = run_qwen_stage(
            run_id,
            manifest,
            database,
            ocr_results,
            template,
            config,
            store,
            extract_meter,
            search_meter,
            qwen_total_meter,
        )
        stage_timings["qwen_ms"] = round((time.perf_counter() - qwen_started) * 1000)
        progress("Qwen 阶段完成，开始校验和导出")
        for product in manifest["products"]:
            key = identity_key(product["platform"], product["product_id"])
            document = assemble_product(
                product,
                database.get(key),
                db_artifact_ref,
                ocr_results[key],
                extract_results.get(key, skipped_qwen("MISSING", "缺少提取结果")),
                search_results.get(key, skipped_qwen("MISSING", "缺少搜索结果")),
                template,
            )
            documents.append(document)
            product_path = products_dir / f"{product['product_id']}.json"
            dump_json(product_path, document)
            price = document.get("database", {}).get("price", {})
            store.event(
                run_id,
                "validation",
                document["status"],
                platform=product["platform"],
                product_id=product["product_id"],
                severity="warning" if document["validation_issues"] else "info",
                message=(
                    "; ".join(item["code"] for item in document["validation_issues"])
                    if document["validation_issues"]
                    else "校验通过"
                ),
                source_table=price.get("source_table"),
                source_column=price.get("source_column"),
                source_time=price.get("source_time"),
                artifact_ref=str(product_path),
            )
        review_images = [
            {
                "platform": product["platform"],
                "product_id": product["product_id"],
                "image_name": item["image_name"],
                "error": item.get("error"),
            }
            for product in manifest["products"]
            for item in ocr_results[identity_key(product["platform"], product["product_id"])]
            if not item.get("ok")
        ]
        dump_json(output_dir / "待复核图片.json", review_images)
        events = store.events(run_id)
        workbook = None
        csv_file = export_csv_file(output_dir, template, documents)
        if not skip_workbook:
            workbook = export_workbook_file(
                output_dir, workbook_payload(run_id, template, documents, events), config
            )
        run_status = "complete" if all(item["status"] == "success" for item in documents) else "review"
        store.finish_run(run_id, run_status)
        peaks = {
            "ocr": ocr_meter.peak,
            "ocr_limit": ocr_worker_count(config),
            "qwen_extract": extract_meter.peak,
            "qwen_search": search_meter.peak,
            "qwen_total": qwen_total_meter.peak,
        }
        ocr_items = [result for results in ocr_results.values() for result in results]
        performance = build_performance_summary(
            stage_timings,
            ocr_items,
            config.get("metrics", {}),
            peaks,
            total_ms=round((time.perf_counter() - pipeline_started) * 1000),
            planned_images=sum(len(product["images"]) for product in manifest["products"]),
        )
        performance.update(
            {
                "run_id": run_id,
                "attempt_id": attempt_id,
                "run_fingerprint": run_fingerprint,
                "generated_at": iso_now(),
                "ocr_provider": config.get("ocr_provider", "vllm-qwen-vl"),
                "ocr_model": config.get("paddle_ocr_api_url") if config.get("ocr_provider") == "paddleocr-vl" else config.get("vllm_ocr_model"),
                "ocr_model_version": config.get("paddle_ocr_model_version") if config.get("ocr_provider") == "paddleocr-vl" else config.get("vllm_ocr_model_version"),
            }
        )
        previous = previous_performance(output_root, output_dir, run_fingerprint)
        if previous:
            performance["comparison_to_previous"] = compare_performance(performance, previous)
        performance_path = output_dir / "performance.json"
        dump_json(performance_path, performance)
        current_events = store.events(run_id, attempt_id)
        report = write_report(
            output_dir,
            run_id,
            run_status,
            documents,
            current_events,
            config.get("metrics", {}),
            peaks,
            performance,
        )
        return {
            "run_id": run_id,
            "attempt_id": attempt_id,
            "status": run_status,
            "resumed": resumed,
            "output_dir": str(output_dir),
            "workbook": str(workbook) if workbook else None,
            "csv": str(csv_file),
            "report": str(report),
            "success_count": sum(item["status"] == "success" for item in documents),
            "review_count": sum(item["status"] == "review" for item in documents),
            "peaks": peaks,
            "metrics": dict(config.get("metrics", {})),
            "performance": performance,
            "performance_path": str(performance_path),
        }
    except DemoError as exc:
        if run_id and exc.code != "SIMULATED_INTERRUPT":
            store.finish_run(
                run_id,
                "interrupted" if exc.code == "PADDLE_OCR_INTERRUPTED" else "failed",
            )
            peaks = {
                "ocr": ocr_meter.peak,
                "ocr_limit": ocr_worker_count(config),
                "qwen_extract": extract_meter.peak,
                "qwen_search": search_meter.peak,
                "qwen_total": qwen_total_meter.peak,
            }
            performance = None
            failure_documents = documents or list(partial_documents.values())
            if isinstance(exc, OcrStageError):
                ocr_items = [
                    result
                    for results in exc.ocr_results.values()
                    for result in results
                ]
                performance = build_performance_summary(
                    {
                        "ocr_ms": exc.duration_ms,
                        "parallel_wall_ms": round((time.perf_counter() - pipeline_started) * 1000),
                    },
                    ocr_items,
                    config.get("metrics", {}),
                    peaks,
                    total_ms=round((time.perf_counter() - pipeline_started) * 1000),
                    planned_images=exc.total_images,
                )
                performance.update(
                    {
                        "run_id": run_id,
                        "attempt_id": attempt_id,
                        "run_fingerprint": run_fingerprint,
                        "generated_at": iso_now(),
                        "ocr_provider": config.get("ocr_provider", "vllm-qwen-vl"),
                        "ocr_model": config.get("paddle_ocr_api_url") if config.get("ocr_provider") == "paddleocr-vl" else config.get("vllm_ocr_model"),
                        "ocr_model_version": config.get("paddle_ocr_model_version") if config.get("ocr_provider") == "paddleocr-vl" else config.get("vllm_ocr_model_version"),
                        "failure": {"code": exc.code, "message": safe_message(exc)},
                    }
                )
                dump_json(output_dir / "performance.json", performance)
            write_report(
                output_dir,
                run_id,
                "failed",
                failure_documents,
                store.events(run_id, attempt_id),
                config.get("metrics", {}),
                peaks,
                performance,
            )
        raise
    finally:
        store.close()


def validate_template(template: dict[str, Any]) -> None:
    if not isinstance(template, dict) or not template.get("version"):
        raise DemoError("INVALID_TEMPLATE", "模板缺少 version")
    if template.get("model", {}).get("name") != "qwen-plus":
        raise DemoError("INVALID_TEMPLATE", "Demo 模型必须保持 qwen-plus")
    model_fields = template.get("model", {}).get("fields", [])
    search_fields = template.get("search", {}).get("fields", [])
    output_columns = template.get("output_columns", [])
    for label, fields in (
        ("model.fields", model_fields),
        ("search.fields", search_fields),
        ("output_columns", output_columns),
    ):
        names = [field.get("name") for field in fields]
        if any(not name for name in names) or len(names) != len(set(names)):
            raise DemoError("INVALID_TEMPLATE", f"{label} 字段名为空或重复")
    if not output_columns:
        raise DemoError("INVALID_TEMPLATE", "output_columns 不能为空")
    protected = set(template.get("source_policy", {}).get("protected_from_model", []))
    if not {"platform", "product_id", "价格"}.issubset(protected):
        raise DemoError("INVALID_TEMPLATE", "模板必须保护 platform、product_id 和价格")
    model_field_names = {str(field.get("name")) for field in model_fields}
    missing_enrichment = set(ENRICHMENT_FIELDS) - model_field_names
    if missing_enrichment:
        raise DemoError(
            "INVALID_TEMPLATE",
            "v2 模板缺少 OCR 明确提取字段: " + ", ".join(sorted(missing_enrichment)),
        )


def create_mock_batch(root: Path, include_failure: bool = False) -> Path:
    product_ids = ["497394", "100002", "100003", "100004", "100005"]
    root.mkdir(parents=True, exist_ok=True)
    for product_id in product_ids:
        folder = root / product_id
        folder.mkdir(parents=True, exist_ok=True)
        for number in (1, 2):
            name = (
                "02_fail_ocr.jpg"
                if include_failure and product_id == "100005" and number == 2
                else f"{number:02}.jpg"
            )
            # Only mock mode reads these fixtures; unique bytes keep SHA-based caching honest.
            (folder / name).write_bytes(
                b"\xff\xd8\xff\xe0" + f"mock:{product_id}:{name}".encode("ascii")
            )
    return root


def assume_yes() -> bool:
    """是否处于无人值守模式（Linux 后台服务 / systemd）。

    服务进程没有 tty，任何 ``input()`` / ``getpass()`` 都会让进程永久挂起，
    所以后台运行时由 ``OCR_ASSUME_YES=1``（或命令行 ``--assume-yes``）打开：
    交互确认一律跳过，缺少的必填项改为直接报错退出，而不是等待输入。
    """
    return os.environ.get("OCR_ASSUME_YES", "").strip().casefold() in {
        "1",
        "true",
        "yes",
        "y",
        "on",
    }


def required_text(env_name: str, label: str, *, secret: bool = False) -> str:
    value = os.environ.get(env_name, "").strip()
    if not value:
        if assume_yes():
            raise DemoError(
                "CONFIG_REQUIRED",
                f"缺少 {label}；无人值守模式不会等待输入，请通过环境变量 {env_name} 提供",
            )
        value = (getpass.getpass(f"{label}: ") if secret else input(f"{label}: ")).strip()
    if not value:
        raise DemoError("CONFIG_REQUIRED", f"缺少 {label}")
    return value


def optional_text(env_name: str, label: str, default: str = "") -> str:
    if env_name in os.environ:
        return os.environ[env_name].strip()
    if assume_yes():
        return default
    suffix = f" [{default}]" if default else "（可留空）"
    value = input(f"{label}{suffix}: ").strip()
    return value or default


def ensure_https(url: str, label: str) -> str:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme.lower() != "https" or not parsed.netloc:
        raise DemoError("TLS_REQUIRED", f"{label} 必须是有效的 https:// 地址")
    return url.rstrip("/")


def ensure_vllm_api_base(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc or not parsed.path.rstrip("/").endswith("/v1"):
        raise DemoError("INVALID_VLLM_URL", "VLLM_OCR_API_BASE 必须是以 /v1 结尾的完整地址")
    if parsed.scheme.lower() == "http":
        try:
            private_address = ipaddress.ip_address(parsed.hostname or "").is_private
        except ValueError:
            private_address = False
        if not private_address:
            raise DemoError("VLLM_TLS_REQUIRED", "非内网 IP 的 VLLM 地址必须使用 https://")
    return url.rstrip("/")


def ensure_paddle_api_url(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc or not parsed.path.rstrip("/").endswith("/v1/ocr"):
        raise DemoError("INVALID_PADDLE_URL", "PADDLE_OCR_API_URL 必须是以 /v1/ocr 结尾的完整地址")
    return url.rstrip("/")


def display_manifest(
    manifest: dict[str, Any],
    *,
    vllm_ocr_url: str,
    vllm_ocr_model: str,
    qwen_url: str,
    redshift_target: str,
    mock: bool,
) -> None:
    print("\n========== 本次处理与外发清单 ==========")
    for product in manifest["products"]:
        print(
            f"平台={product['platform']}  product_id={product['product_id']}  "
            f"图片={len(product['images'])}"
        )
        for image in product["images"]:
            print(f"  - {image['path']}  sha256={image['sha256']}")
    if mock:
        print("\n目标：本地模拟服务（不会外发图片、OCR 文本或商品数据）")
    else:
        print(f"\n本地 Qwen-VL OCR：{vllm_ocr_url}/chat/completions，模型={vllm_ocr_model}（发送所列图片）")
        print(f"阿里 Qwen：{qwen_url}/chat/completions（发送商品名、OCR 文本和属性；搜索节点会联网）")
        print(f"PostgreSQL：{redshift_target}（发送所列 platform/product_id 查询）")
    print("========================================\n")


def postgres_connection_info(
    *,
    host: str,
    port: int,
    database: str,
    user: str,
    password: str,
    sslmode: str,
) -> dict[str, Any]:
    return {
        "host": host,
        "port": port,
        "dbname": database,
        "user": user,
        "password": password,
        "client_encoding": "UTF8",
        "sslmode": sslmode,
        "connect_timeout": 20,
        "options": "-c statement_timeout=60000",
        "application_name": "windows_ocr_demo",
    }


def real_config(manifest: dict[str, Any]) -> dict[str, Any]:
    paddle_ocr_url = ensure_paddle_api_url(os.environ.get("PADDLE_OCR_API_URL", PADDLE_DEFAULT_API_URL).strip())
    paddle_ocr_model_version = os.environ.get("PADDLE_OCR_MODEL_VERSION", PADDLE_DEFAULT_MODEL_VERSION).strip() or PADDLE_DEFAULT_MODEL_VERSION
    qwen_url = ensure_https(os.environ.get("DASHSCOPE_BASE_URL", QWEN_BASE_URL).strip(), "DashScope 地址")
    host = required_text("POSTGRES_HOST", "PostgreSQL 主机")
    port_text = os.environ.get("POSTGRES_PORT", "5432").strip()
    try:
        port = int(port_text)
        if not 1 <= port <= 65535:
            raise ValueError
    except ValueError as exc:
        raise DemoError("INVALID_PORT", f"PostgreSQL 端口无效: {port_text}") from exc
    database = required_text("POSTGRES_DATABASE", "PostgreSQL 数据库名")
    user = required_text("POSTGRES_USER", "PostgreSQL 只读用户名")
    schema = os.environ.get("POSTGRES_SCHEMA", "").strip()
    sslmode = os.environ.get("POSTGRES_SSLMODE", "verify-full").strip().lower()
    if sslmode not in {"disable", "require", "verify-ca", "verify-full"}:
        raise DemoError(
            "INVALID_SSLMODE",
            "POSTGRES_SSLMODE 只能是 disable、require、verify-ca 或 verify-full",
        )
    display_manifest(
        manifest,
        vllm_ocr_url=paddle_ocr_url,
        vllm_ocr_model=paddle_ocr_model_version,
        qwen_url=qwen_url,
        redshift_target=f"{host}:{port}/{database}",
        mock=False,
    )
    if assume_yes():
        progress("[无人值守] 已自动确认数据外发清单（OCR_ASSUME_YES=1）")
    elif input("确认允许发送以上数据请输入 SEND: ").strip() != "SEND":
        raise DemoError("CANCELLED", "用户取消外发")
    rotated = os.environ.get("OCR_DEMO_KEYS_ROTATED", "").strip().upper()
    if rotated not in {"YES", "ROTATED"}:
        if assume_yes():
            raise DemoError(
                "KEY_ROTATION_REQUIRED",
                "无人值守模式必须通过环境变量 OCR_DEMO_KEYS_ROTATED=YES 显式确认密钥已轮换",
            )
        rotated = input("确认此前暴露的阿里密钥已吊销并轮换，请输入 ROTATED: ").strip().upper()
    if rotated != "ROTATED" and rotated != "YES":
        raise DemoError("KEY_ROTATION_REQUIRED", "未确认密钥轮换，禁止真实调用")
    qwen_key = required_text("DASHSCOPE_API_KEY", "已轮换的 DashScope API Key（隐藏输入）", secret=True)
    password = required_text("POSTGRES_PASSWORD", "PostgreSQL 密码（隐藏输入）", secret=True)
    return {
        "mock": False,
        "ocr_provider": "paddleocr-vl",
        "paddle_ocr_api_url": paddle_ocr_url,
        "paddle_ocr_model_version": paddle_ocr_model_version,
        "qwen_base_url": qwen_url,
        "qwen_api_key": qwen_key,
        "redshift_connection": postgres_connection_info(
            host=host,
            port=port,
            database=database,
            user=user,
            password=password,
            sslmode=sslmode,
        ),
        "redshift_schema": schema,
        "metrics": {},
        "metrics_lock": threading.Lock(),
    }


def mock_config() -> dict[str, Any]:
    return {
        "mock": True,
        "vllm_ocr_api_base": "http://192.168.1.115:8801/v1",
        "vllm_ocr_api_key": "EMPTY",
        "vllm_ocr_model": "Qwen/Qwen3.8-27B",
        "vllm_ocr_model_version": "mock-v1",
        "qwen_base_url": "mock://qwen",
        "qwen_api_key": "",
        "ocr_retry_delays": (0.0, 0.0, 0.0, 0.0),
        "qwen_retry_delays": (0.0, 0.0, 0.0, 0.0),
        "redshift_connection": {},
        "redshift_schema": None,
        "metrics": {},
        "metrics_lock": threading.Lock(),
    }


def prepare_manifest(args: argparse.Namespace, state_dir: Path) -> dict[str, Any]:
    if args.mock and not args.root:
        root = create_mock_batch(state_dir / "mock_batch")
    else:
        root = Path(args.root).resolve() if args.root else choose_root()
    if not root.is_dir():
        raise DemoError("ROOT_NOT_FOUND", f"批次目录不存在: {root}")
    candidates = product_candidates(root)
    if args.products:
        selected = choose_products(root, args.products)
    elif args.mock and 1 <= len(candidates) <= MAX_PRODUCTS:
        selected = candidates
    else:
        selected = choose_products(root, None)
    platform = (args.platform or os.environ.get("OCR_DEMO_PLATFORM", "")).strip()
    if not platform:
        if assume_yes():
            raise DemoError(
                "PLATFORM_REQUIRED",
                "无人值守模式必须通过 --platform 或 OCR_DEMO_PLATFORM 指定数据库平台标识",
            )
        suggested = root.parent.name
        entered = input(f"请输入数据库平台标识（建议值 {suggested!r}，回车确认）: ").strip()
        platform = entered or suggested
    return build_manifest(root, selected, platform)


def open_outputs(result: dict[str, Any]) -> None:
    if os.name != "nt" or not hasattr(os, "startfile"):
        return
    for target in (result.get("report"), result.get("output_dir")):
        if target:
            try:
                os.startfile(target)  # type: ignore[attr-defined]
            except OSError:
                pass


def self_test() -> dict[str, Any]:
    assert parse_spec("60粒/瓶") == {
        "main_spec": "60粒", "package": "1瓶", "total": 60, "unit": "粒", "status": "parsed"
    }
    assert parse_spec("30条*3盒")["total"] == 90
    assert parse_spec("2盒×30粒")["total"] == 60
    assert parse_spec("500mg*60粒")["total"] == 60
    assert parse_spec("100片*3礼盒装")["total"] == 300
    assert parse_spec("900g*2")["total"] == 1800
    assert parse_spec("80支/礼盒装")["package"] == "1礼盒装"
    drop_spec = parse_spec("每瓶90滴（2.5ml）")
    assert drop_spec["total"] == 90 and drop_spec["unit"] == "滴"
    costs = calculate_costs("37", 60, 2, 2, "粒", "每日2粒")
    assert costs["unit_price"] == Decimal("0.6167")
    assert costs["min_daily_cost"] == Decimal("1.2333")
    calcium_costs = calculate_costs("147.05", 300, 2, 2, "片", "每日2粒")
    assert calcium_costs["unit_price"] == Decimal("0.4902")
    assert calcium_costs["min_daily_cost"] == Decimal("0.9803")
    tea_costs = calculate_costs("178", 60, 1, 1, "袋", "饭后半小时一包")
    assert tea_costs["min_daily_cost"] == Decimal("2.9667")
    drop_costs = calculate_costs("108.75", 90, 1, 1, "滴", "每日1滴")
    assert drop_costs["min_daily_cost"] == Decimal("1.2083")
    assert comparable_brand("雀巢健康科学（NESTLE HEALTH SCIENCE）") == comparable_brand("雀巢健康科学")
    retry_count = 0

    def transient_call() -> str:
        nonlocal retry_count
        retry_count += 1
        if retry_count < 3:
            raise DemoError("NETWORK_ERROR", "temporary", retryable=True)
        return "ok"

    retry_value, retry_attempts = call_with_retry(
        transient_call, max_attempts=3, retry_delays=(0.0, 0.0)
    )
    assert retry_value == "ok" and retry_attempts == 3
    assert blue_hat_value({"批准文号": "国食健字G20040371"}, None)[0] == "国食健字G20040371"
    assert blue_hat_value({"批准文号": "国食注字TY20230052"}, None) == (
        None,
        "国食注字TY20230052",
    )
    template = load_json(Path(__file__).with_name("template-v2.json"))
    validate_template(template)
    smaller_template = json.loads(json.dumps(template, ensure_ascii=False))
    smaller_template["output_columns"] = smaller_template["output_columns"][:-1]
    validate_template(smaller_template)
    valid = mock_qwen_values(template["model"]["fields"], False)
    assert not validate_object(valid, template["model"]["fields"])
    assert validate_object({**valid, "price": 37}, template["model"]["fields"])
    assert identity_key("jd", "1") != identity_key("tmall", "1")
    assert stable_hash({"a": 1, "b": 2}) == stable_hash({"b": 2, "a": 1})
    assert provider_search_sources(
        {"search_info": {"search_results": [{"title": "证据", "url": "https://example.com/a"}]}}
    ) == [{"url": "https://example.com/a", "title": "证据"}]

    evidence = [{"url": "https://example.com/factory", "title": "生产信息"}]
    ocr_priority = resolve_enrichment(
        {"factory_name": "OCR生产商"},
        {"生产企业": "数据库生产商"},
        {"factory_name": "搜索生产商"},
        evidence,
    )
    assert ocr_priority["values"]["factory_name"] == "OCR生产商"
    assert ocr_priority["selected_sources"]["factory_name"]["source"] == "qwen_extract"
    database_priority = resolve_enrichment(
        {"factory_name": None},
        {"生产企业": "数据库生产商"},
        {"factory_name": "搜索生产商"},
        evidence,
    )
    assert database_priority["values"]["factory_name"] == "数据库生产商"
    assert database_priority["selected_sources"]["factory_name"]["attribute_name"] == "生产企业"
    search_fallback = resolve_enrichment(
        {}, {}, {"factory_name": "搜索生产商"}, evidence
    )
    assert search_fallback["values"]["factory_name"] == "搜索生产商"
    assert search_fallback["selected_sources"]["factory_name"]["source"] == "qwen_search"
    rejected_search = resolve_enrichment(
        {}, {}, {"factory_name": "无URL的搜索生产商"}, []
    )
    assert rejected_search["values"]["factory_name"] is None
    assert rejected_search["search_rejected"] is True
    complete_attributes = {
        "生产企业": "数据库生产商",
        "生产地址": "数据库地址",
        "食品生产许可证编号": "SC12345678901234",
        "产地": "中国",
    }
    assert missing_search_fields({}, complete_attributes, include_hot_topics=False) == []
    assert missing_search_fields({}, {}, include_hot_topics=False) == [
        "factory_name",
        "factory_address",
        "production_license",
        "origin",
    ]
    assert missing_search_fields({}, complete_attributes, include_hot_topics=True) == ["hot_topics"]

    def qwen_search_test(
        product_id: str,
        enrichment_attributes: dict[str, str],
        *,
        bulk_first_pass: bool = False,
    ) -> tuple[int, dict[str, Any]]:
        product = {"platform": "jd", "product_id": product_id}
        record = next(iter(mock_database_results({"products": [product]}).values()))
        record["attributes"] = [
            *record["attributes"],
            *(
                {"name": name, "value": value, "source_table": "mock_attributes"}
                for name, value in enrichment_attributes.items()
            ),
        ]
        record["db_snapshot_hash"] = stable_hash(record)
        key = identity_key("jd", product_id)
        config = mock_config()
        config["include_hot_topics"] = False
        config["bulk_first_pass"] = bulk_first_pass
        with tempfile.TemporaryDirectory(prefix="ocr-demo-qwen-stage-") as temporary:
            store = StateStore(Path(temporary))
            try:
                _, searches = run_qwen_stage(
                    "qwen-stage-test",
                    {"products": [product]},
                    {key: record},
                    {
                        key: [
                            {
                                "ok": True,
                                "image_name": "01.jpg",
                                "markdown": "模拟OCR文本",
                                "sha256": "mock-sha256",
                            }
                        ]
                    },
                    template,
                    config,
                    store,
                    ConcurrencyMeter(),
                    ConcurrencyMeter(),
                    ConcurrencyMeter(),
                )
            finally:
                store.close()
        return config["metrics"].get("qwen_search_api_calls", 0), searches[key]

    complete_calls, complete_search = qwen_search_test("complete", complete_attributes)
    assert complete_calls == 0
    assert complete_search.get("skipped") is True
    missing_calls, missing_search = qwen_search_test("missing", {})
    assert missing_calls == 1 and missing_search.get("ok") is True
    bulk_calls, bulk_search = qwen_search_test("bulk", {}, bulk_first_pass=True)
    assert bulk_calls == 0
    assert bulk_search.get("skipped") is True

    with tempfile.TemporaryDirectory() as temporary:
        batch_root = Path(temporary)
        product_dirs = []
        for index in range(MAX_PRODUCTS + 1):
            product_dir = batch_root / str(100000 + index)
            product_dir.mkdir()
            (product_dir / "01.jpg").write_bytes(str(index).encode())
            product_dirs.append(product_dir)
        assert len(build_manifest(batch_root, product_dirs[:MAX_PRODUCTS], "jd")["products"]) == MAX_PRODUCTS
        try:
            build_manifest(batch_root, product_dirs, "jd")
        except DemoError as exc:
            assert exc.code == "PRODUCT_COUNT"
        else:
            raise AssertionError("超过100个商品时必须拒绝运行")

    conflict_product = {"platform": "jd", "product_id": "1242256"}
    conflict_record = next(iter(mock_database_results({"products": [conflict_product]}).values()))
    conflict_model = mock_qwen_values(template["model"]["fields"], False)
    conflict_model.update({"guige": "100片", "guige_zong_liang": 300, "is_multi_spec": True})
    conflict_document = assemble_product(
        conflict_product,
        conflict_record,
        "mock-db.json",
        [],
        {"ok": True, "data": conflict_model, "raw_json_ref": "mock-extract.json"},
        {
            "ok": True,
            "data": {"factory_name": "无证据的模型内容"},
            "provider_sources": [],
            "raw_json_ref": "mock-search.json",
        },
        template,
    )
    assert conflict_document["status"] == "review"
    assert any(item["code"] == "SPEC_TOTAL_CONFLICT" for item in conflict_document["validation_issues"])
    assert any(item["code"] == "SEARCH_EVIDENCE_MISSING" for item in conflict_document["validation_issues"])
    assert conflict_document["fields"]["代工厂"] is None

    warning_model = mock_qwen_values(template["model"]["fields"], False)
    warning_model.update(
        {"guige": "100片*3礼盒装", "guige_zong_liang": 300, "is_multi_spec": True}
    )
    warning_document = assemble_product(
        conflict_product,
        conflict_record,
        "mock-db.json",
        [],
        {"ok": True, "data": warning_model, "raw_json_ref": "mock-extract.json"},
        {
            "ok": True,
            "data": {"factory_name": "无证据的模型内容"},
            "provider_sources": [],
            "raw_json_ref": "mock-search.json",
        },
        template,
    )
    assert warning_document["status"] == "success"
    assert next(
        item for item in warning_document["validation_issues"]
        if item["code"] == "SEARCH_EVIDENCE_MISSING"
    )["severity"] == "warning"
    assert warning_document["fields"]["代工厂"] is None
    warning_payload = workbook_payload("warning-test", template, [warning_document], [])
    assert len(warning_payload["result"]["rows"]) == 1
    assert warning_payload["result"]["rows"][0]["商品主数据来源"] == "京东（主平台，platform_key=1）"
    assert warning_payload["result"]["rows"][0]["价格数据来源"] == "京东（主平台，platform_key=1）"

    packaging_model = mock_qwen_values(template["model"]["fields"], False)
    packaging_model.update({"guige": "400g", "guige_zong_liang": 400})
    stale_record = {
        **conflict_record,
        "price": {
            **conflict_record["price"],
            "source_table": "mv_com_goods_statistics_monthly_v2_internal_ssv4",
            "source_time": "2000-01-01",
        },
        "attributes": [
            *conflict_record["attributes"],
            {"name": "包装清单", "value": "蔼儿舒400g*1"},
            {"name": "包装规格", "value": "1罐"},
            {"name": "包装形式", "value": "罐装"},
        ],
    }
    stale_document = assemble_product(
        conflict_product,
        stale_record,
        "mock-db.json",
        [],
        {"ok": True, "data": packaging_model, "raw_json_ref": "mock-extract.json"},
        {"ok": True, "data": {}, "provider_sources": [], "raw_json_ref": None},
        template,
    )
    assert stale_document["fields"]["包装"] == "1罐"
    assert next(
        item for item in stale_document["validation_issues"] if item["code"] == "PRICE_STALE"
    )["severity"] == "warning"

    meter = ConcurrencyMeter()
    def metered() -> None:
        with meter:
            time.sleep(0.02)
    with futures.ThreadPoolExecutor(max_workers=5) as executor:
        list(executor.map(lambda _: metered(), range(12)))
    assert meter.peak <= 5

    import redshift_backend
    redshift_backend._self_test()

    with tempfile.TemporaryDirectory(prefix="ocr-demo-self-test-") as temporary:
        base = Path(temporary)
        state = base / "state"
        output = base / "runs"
        cache_store = StateStore(base / "cache-metadata")
        first_image_path = base / "first.jpg"
        second_image_path = base / "renamed.jpg"
        first_image_path.write_bytes(b"same-image")
        second_image_path.write_bytes(b"same-image")
        shared_sha = file_sha256(first_image_path)
        cache_product = {"platform": "jd", "product_id": "cache-test"}
        cache_config = mock_config()
        first_cached = run_ocr_one(
            "cache-test-run",
            cache_product,
            {"name": first_image_path.name, "path": str(first_image_path), "sha256": shared_sha},
            cache_config,
            cache_store,
            ConcurrencyMeter(),
        )
        second_cached = run_ocr_one(
            "cache-test-run",
            cache_product,
            {"name": second_image_path.name, "path": str(second_image_path), "sha256": shared_sha},
            cache_config,
            cache_store,
            ConcurrencyMeter(),
        )
        cache_store.close()
        assert not first_cached["cached"] and second_cached["cached"]
        assert second_cached["image_name"] == "renamed.jpg"
        assert second_cached["image_path"] == str(second_image_path)
        root = create_mock_batch(base / "batch")
        manifest = build_manifest(root, product_candidates(root), "jd")
        config = mock_config()
        try:
            execute_pipeline(
                manifest,
                Path(__file__).with_name("template-v2.json"),
                state,
                output,
                config,
                interrupt_after_parallel=True,
                skip_workbook=True,
            )
            raise AssertionError("simulated interruption did not stop")
        except DemoError as exc:
            assert exc.code == "SIMULATED_INTERRUPT"
        first_ocr_calls = config["metrics"].get("ocr_api_calls")
        assert first_ocr_calls == 10
        result = execute_pipeline(
            manifest,
            Path(__file__).with_name("template-v2.json"),
            state,
            output,
            config,
        )
        assert result["resumed"] is True
        assert result["attempt_id"].startswith("attempt-")
        assert result["success_count"] == 5 and result["review_count"] == 0
        assert config["metrics"].get("ocr_api_calls") == first_ocr_calls
        assert result["peaks"]["ocr"] <= OCR_WORKERS
        assert result["peaks"]["qwen_extract"] <= QWEN_EXTRACT_WORKERS
        assert result["peaks"]["qwen_search"] <= QWEN_SEARCH_WORKERS
        assert result["peaks"]["qwen_total"] <= QWEN_EXTRACT_WORKERS + QWEN_SEARCH_WORKERS
        for required in (result["workbook"], result["report"]):
            assert required and Path(required).is_file() and Path(required).stat().st_size > 0
        report_text = Path(result["report"]).read_text(encoding="utf-8")
        assert result["attempt_id"] in report_text and "https://example.com/mock-source" in report_text
        product_json = load_json(Path(result["output_dir"]) / "products" / "497394.json")
        assert product_json["database"]["key"]["value"] == 22129910
        assert as_decimal(product_json["fields"]["价格"]) == Decimal("103.90")
        assert as_decimal(product_json["fields"]["价格"]) != Decimal("37")
        bulk_config = mock_config()
        bulk_config["bulk_first_pass"] = True
        bulk_config["include_hot_topics"] = False
        bulk_result = execute_pipeline(
            manifest,
            Path(__file__).with_name("template-v2.json"),
            state,
            output,
            bulk_config,
            skip_workbook=True,
        )
        assert bulk_config["metrics"].get("qwen_search_api_calls", 0) == 0
        bulk_product = load_json(
            Path(bulk_result["output_dir"]) / "products" / "497394.json"
        )
        assert bulk_product["missing_enrichment_fields"] == [
            "factory_name",
            "factory_address",
            "production_license",
            "origin",
            "hot_topics",
        ]
        connection = sqlite3.connect(state / "state.sqlite3")
        try:
            attempts = connection.execute(
                "SELECT COUNT(DISTINCT attempt_id) FROM events WHERE run_id=? AND attempt_id IS NOT NULL",
                (result["run_id"],),
            ).fetchone()[0]
        finally:
            connection.close()
        assert attempts == 2

        failure_root = create_mock_batch(base / "failure_batch", include_failure=True)
        failure_manifest = build_manifest(failure_root, product_candidates(failure_root), "jd")
        failure_config = mock_config()
        failure_result = execute_pipeline(
            failure_manifest,
            Path(__file__).with_name("template-v2.json"),
            state,
            output,
            failure_config,
            skip_workbook=True,
        )
        assert failure_result["success_count"] == 4 and failure_result["review_count"] == 1
        failed_document = load_json(
            Path(failure_result["output_dir"]) / "products" / "100005.json"
        )
        assert failed_document["status"] == "review"
        assert any(item["code"] == "MOCK_OCR_FAILURE" for item in failed_document["validation_issues"])
        assert failed_document["artifacts"]["qwen_extract"]
        assert failed_document["fields"]["规格"] == "60粒"
    return {
        "ok": True,
        "checks": [
            "specification",
            "real_product_regressions",
            "decimal_cost",
            "drop_unit_cost",
            "json_types",
            "cache_key",
            "retry_policy",
            "hundred_product_manifest",
            "redshift_rules",
            "interrupt_resume",
            "attempt_audit",
            "ocr_cache",
            "ocr_cache_current_metadata",
            "search_without_evidence_warning",
            "ocr_database_search_precedence",
            "conditional_search",
            "bulk_first_pass",
            "price_stale_warning",
            "packaging_priority",
            "failure_review",
            "partial_ocr_draft",
            "concurrency_limits",
            "excel_json_html",
        ],
    }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    app_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="本地 Qwen-VL OCR 工作流 Demo v5")
    parser.add_argument("--root", help="批次根目录（其下为 product_id 目录）")
    parser.add_argument(
        "--products",
        nargs="+",
        metavar="PRODUCT_ID",
        help=f"直接指定 1 到 {MAX_PRODUCTS} 个商品目录名",
    )
    parser.add_argument("--platform", help="数据库平台标识；未给出时交互确认")
    parser.add_argument("--template", default=str(app_dir / "template-v2.json"))
    parser.add_argument("--state-dir", default=str(app_dir / ".state"))
    parser.add_argument("--output-root", default=str(app_dir / "runs"))
    parser.add_argument("--mock", action="store_true", help="使用本地模拟服务，不外发数据")
    parser.add_argument("--self-test", action="store_true", help="运行离线自检")
    parser.add_argument("--new-run", action="store_true", help="不恢复未完成运行，创建新运行")
    parser.add_argument(
        "--ocr-workers",
        type=int,
        choices=range(1, OCR_MAX_WORKERS + 1),
        default=OCR_WORKERS,
        help=f"本地 OCR 并发数（1 到 {OCR_MAX_WORKERS}，默认 {OCR_WORKERS}）",
    )
    parser.add_argument(
        "--force-ocr",
        action="store_true",
        help="仅用于测速：忽略 OCR 成功缓存，强制重新请求本地 OCR",
    )
    parser.add_argument(
        "--bulk-first-pass",
        action="store_true",
        help="批量首轮：跳过联网搜索，只输出 OCR、数据库结果和待补全字段",
    )
    parser.add_argument("--open", action="store_true", help="成功后打开报告和输出目录")
    parser.add_argument(
        "--assume-yes",
        action="store_true",
        help="无人值守：跳过所有交互确认（等价于环境变量 OCR_ASSUME_YES=1），供后台服务使用",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.assume_yes:
        os.environ["OCR_ASSUME_YES"] = "1"
    if args.self_test:
        result = self_test()
        print("SELF_TEST_OK=" + json.dumps(result, ensure_ascii=False))
        return 0
    template_path = Path(args.template).resolve()
    template = load_json(template_path)
    validate_template(template)
    state_dir = Path(args.state_dir).resolve()
    output_root = Path(args.output_root).resolve()
    manifest = prepare_manifest(args, state_dir)
    if args.mock:
        config = mock_config()
        display_manifest(
            manifest,
            vllm_ocr_url=config["vllm_ocr_api_base"],
            vllm_ocr_model=config["vllm_ocr_model"],
            qwen_url=config["qwen_base_url"],
            redshift_target="mock://redshift",
            mock=True,
        )
    else:
        config = real_config(manifest)
    config["ocr_workers"] = args.ocr_workers
    config["paddle_workers"] = min(2, args.ocr_workers)
    config["force_ocr"] = args.force_ocr
    config["bulk_first_pass"] = args.bulk_first_pass
    config["include_hot_topics"] = not args.bulk_first_pass
    result = execute_pipeline(
        manifest,
        template_path,
        state_dir,
        output_root,
        config,
        force_new=args.new_run,
    )
    print(f"OUTPUT_DIR={result['output_dir']}")
    if result.get("workbook"):
        print(f"WORKBOOK={result['workbook']}")
    print(f"REPORT={result['report']}")
    print(f"STATUS={result['status']} SUCCESS={result['success_count']} REVIEW={result['review_count']}")
    if args.open:
        open_outputs(result)
    return 0 if result["status"] in {"complete", "review"} else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n已中断；下次使用相同批次会从 SQLite 状态恢复。", file=sys.stderr)
        raise SystemExit(130)
    except DemoError as exc:
        print(f"错误 [{exc.code}]：{safe_message(exc)}", file=sys.stderr)
        raise SystemExit(2)
    except Exception as exc:
        print(f"未处理错误：{safe_message(exc)}", file=sys.stderr)
        raise SystemExit(1)
